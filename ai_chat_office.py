from __future__ import annotations

# Loaded into the parent module's global namespace by the parent file.
# Keep this file focused on the extracted feature area.

def text_attachment_preview(target: str) -> str:
    path = Path(str(target or ""))
    try:
        if not path.is_file() or path.suffix.lower() not in TEXT_ATTACHMENT_EXTS:
            return ""
        data = path.read_bytes()
    except OSError:
        return ""
    truncated = len(data) > TEXT_ATTACHMENT_MAX_BYTES
    data = data[:TEXT_ATTACHMENT_MAX_BYTES]
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "gbk", "utf-16", "big5"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = data.decode("utf-8", errors="replace")
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return ""
    excerpt = text[:TEXT_ATTACHMENT_PREVIEW_CHARS]
    suffix = "…" if truncated or len(text) > TEXT_ATTACHMENT_PREVIEW_CHARS else ""
    return f"{excerpt}{suffix}"


def office_attachment_preview(target: str, max_chars: int = OFFICE_ATTACHMENT_PREVIEW_CHARS) -> str:
    path = Path(str(target or ""))
    suffix = path.suffix.lower()
    try:
        if not path.is_file() or suffix not in OFFICE_ATTACHMENT_EXTS:
            return ""
        if suffix == ".docx":
            text = _extract_docx_text(path)
        elif suffix in (".xlsx", ".xlsm"):
            text = _extract_xlsx_text(path)
        elif suffix in (".pptx", ".pptm"):
            text = _extract_pptx_text(path)
        else:
            text = ""
    except Exception as exc:
        return f"Office 预览读取失败：{exc}"
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return "未提取到可读文本。"
    truncated = len(text) > max_chars
    return text[:max_chars] + ("…" if truncated else "")


def office_edit_copy(target: str, operation: str, **kwargs) -> str:
    """Edit an Office file by creating a new copy; never overwrites the source."""
    source = Path(str(target or "")).expanduser()
    if not source.is_file():
        raise ValueError(f"Office 文件不存在：{source}")
    suffix = source.suffix.lower()
    if suffix not in OFFICE_ATTACHMENT_EXTS:
        raise ValueError("仅支持 .docx/.xlsx/.xlsm/.pptx/.pptm")
    op = str(operation or "").strip().lower()
    destination = _office_edit_destination(source)
    if suffix == ".docx":
        summary = _edit_docx_copy(source, destination, op, kwargs)
    elif suffix in (".pptx", ".pptm"):
        summary = _edit_pptx_copy(source, destination, op, kwargs)
    else:
        summary = _edit_xlsx_copy(source, destination, op, kwargs)
    return f"{summary}\n已生成编辑副本：{destination}"


def _office_edit_destination(source: Path) -> Path:
    out_dir = AI_DATA_DIR / "OfficeEdits"
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", source.stem).strip("._") or "office"
    candidate = out_dir / f"{safe_stem}_edited_{stamp}{source.suffix}"
    index = 2
    while candidate.exists():
        candidate = out_dir / f"{safe_stem}_edited_{stamp}_{index}{source.suffix}"
        index += 1
    return candidate


def _edit_docx_copy(source: Path, destination: Path, op: str, kwargs: dict) -> str:
    if op in ("replace", "replace_text", "text_replace"):
        find = str(kwargs.get("find") or kwargs.get("search") or "")
        replace = str(kwargs.get("replace") or kwargs.get("text") or "")
        if not find:
            raise ValueError("replace_text 需要 find")
        count = 0
        with zipfile.ZipFile(source) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
        root = _xml_text(xml)
        ns_uri = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        for node in root.findall(f".//{{{ns_uri}}}t"):
            if node.text and find in node.text:
                count += node.text.count(find)
                node.text = node.text.replace(find, replace)
        if count == 0:
            raise ValueError(f"未找到文本：{find}")
        _write_zip_copy(source, destination, {"word/document.xml": _xml_bytes(root)})
        return f"Word 文本替换完成：{count} 处"
    if op in ("append", "append_text", "append_paragraph"):
        text = str(kwargs.get("text") or kwargs.get("content") or "").strip()
        if not text:
            raise ValueError("append_text 需要 text")
        with zipfile.ZipFile(source) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
        insert = (
            "<w:p><w:r><w:t>"
            + html_lib.escape(text)
            + "</w:t></w:r></w:p>"
        )
        xml = xml.replace("</w:body>", insert + "</w:body>")
        _write_zip_copy(source, destination, {"word/document.xml": xml.encode("utf-8")})
        return "Word 段落追加完成"
    # 其余结构化/排版操作统一交给 python-docx 处理（未知操作在那里报错）。
    return _docx_python_edit(source, destination, op, kwargs)


def _docx_align(paragraph, value) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
        "centre": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "居左": WD_ALIGN_PARAGRAPH.LEFT,
        "居中": WD_ALIGN_PARAGRAPH.CENTER, "居右": WD_ALIGN_PARAGRAPH.RIGHT,
        "两端对齐": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    alignment = mapping.get(str(value).strip().lower())
    if alignment is not None:
        paragraph.alignment = alignment


def _docx_apply_run_format(run, kwargs: dict) -> None:
    from docx.shared import Pt, RGBColor  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    size = kwargs.get("size") or kwargs.get("font_size")
    if size:
        run.font.size = Pt(_to_int(size, 12))
    if kwargs.get("bold") is not None:
        run.font.bold = bool(kwargs.get("bold"))
    if kwargs.get("italic") is not None:
        run.font.italic = bool(kwargs.get("italic"))
    if kwargs.get("underline") is not None:
        run.font.underline = bool(kwargs.get("underline"))
    color = _norm_rgb(kwargs.get("color") or kwargs.get("font_color"))
    if color:
        run.font.color.rgb = RGBColor(*color)
    font_name = kwargs.get("font") or kwargs.get("font_name")
    if font_name:
        run.font.name = str(font_name)
        try:  # 中文需同时设 w:eastAsia 才生效。
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), str(font_name))
        except Exception:
            pass


# --- Word 交叉引用等高级域（书签 / REF / PAGEREF / SEQ 题注）---------------------
def _docx_safe_bookmark_name(name: str) -> str:
    """归一书签名：Word 书签名须以字母开头、仅含字母/数字/下划线（中文亦可，但空格/
    标点不行）。把非法字符换成下划线，必要时加前缀，保证同一名字在「定义」与「引用」
    两处归一结果一致。"""
    raw = str(name or "").strip()
    cleaned = re.sub(r"[^0-9A-Za-z_一-鿿]", "_", raw).strip("_")
    if not cleaned:
        cleaned = "ref_" + hashlib.sha1(raw.encode("utf-8", "surrogatepass")).hexdigest()[:8]
    if cleaned[0].isdigit():
        cleaned = "_" + cleaned
    return cleaned[:40]


def _docx_add_field(paragraph, instr_text: str, placeholder: str = "") -> None:
    """在段落末尾追加一个 Word 域（fldChar begin/instrText/separate/占位/end）。"""
    from docx.oxml.ns import qn  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instr_text
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    holder = OxmlElement("w:t"); holder.set(qn("xml:space"), "preserve"); holder.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, holder, end):
        run._r.append(node)


def _docx_enable_update_fields(document) -> None:
    """标记文档「打开时更新所有域」，使 TOC/交叉引用/题注编号无需手动按 F9 即生效。"""
    from docx.oxml.ns import qn  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    try:
        settings = document.settings.element
        if settings.find(qn("w:updateFields")) is None:
            el = OxmlElement("w:updateFields")
            el.set(qn("w:val"), "true")
            settings.append(el)
    except Exception:
        pass


def _docx_next_bookmark_id(document) -> int:
    from docx.oxml.ns import qn  # type: ignore
    ids = []
    for el in document.element.body.iter(qn("w:bookmarkStart")):
        try:
            ids.append(int(el.get(qn("w:id"))))
        except (TypeError, ValueError):
            continue
    return (max(ids) + 1) if ids else 0


def _docx_wrap_bookmark(paragraph, name: str, document) -> None:
    """用 bookmarkStart/End 把整段包成一个书签，作为交叉引用的目标锚点。"""
    from docx.oxml.ns import qn  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore
    bid = str(_docx_next_bookmark_id(document))
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), bid); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), bid)
    p = paragraph._p
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)   # 书签须在 pPr 之后，否则结构非法
    else:
        p.insert(0, start)
    p.append(end)


def _docx_find_paragraph(document, needle: str):
    needle = str(needle or "").strip()
    if not needle:
        return None
    for paragraph in document.paragraphs:
        if needle in paragraph.text:
            return paragraph
    return None


def _docx_python_edit(source: Path, destination: Path, op: str, kwargs: dict) -> str:
    try:
        import docx  # type: ignore
        from docx.shared import Inches  # type: ignore
    except ImportError as exc:
        raise RuntimeError("该 Word 操作需要 python-docx 库") from exc
    document = docx.Document(str(source))
    if op in ("insert_heading", "add_heading", "heading", "标题"):
        text = str(kwargs.get("text") or kwargs.get("heading") or kwargs.get("content") or "").strip()
        if not text:
            raise ValueError("insert_heading 需要 text")
        try:
            level = int(kwargs.get("level", 1) or 1)
        except (TypeError, ValueError):
            level = 1
        heading = document.add_heading(text, level=max(0, min(level, 9)))
        align = kwargs.get("align") or kwargs.get("alignment")
        if align:
            _docx_align(heading, align)
        summary = f"Word 已追加标题（级别 {level}）"
    elif op in ("add_paragraph", "insert_paragraph", "paragraph", "段落"):
        text = str(kwargs.get("text") or kwargs.get("content") or "")
        style = kwargs.get("style")
        try:
            paragraph = document.add_paragraph(style=str(style)) if style else document.add_paragraph()
        except KeyError:
            paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        _docx_apply_run_format(run, kwargs)
        align = kwargs.get("align") or kwargs.get("alignment")
        if align:
            _docx_align(paragraph, align)
        summary = "Word 已追加段落"
    elif op in ("add_bullets", "add_list", "bullets", "list", "项目符号", "列表"):
        items = _coerce_bullets(kwargs.get("items") or kwargs.get("bullets")
                                or kwargs.get("content") or kwargs.get("text"))
        if not items:
            raise ValueError("add_bullets 需要 items（列表项）")
        ordered = bool(kwargs.get("ordered") or kwargs.get("numbered"))
        style = "List Number" if ordered else "List Bullet"
        for item in items:
            try:
                document.add_paragraph(str(item), style=style)
            except KeyError:
                document.add_paragraph(("" if ordered else "• ") + str(item))
        summary = f"Word 已追加{'有序' if ordered else '项目符号'}列表（{len(items)} 项）"
    elif op in ("insert_table", "add_table", "table", "表格"):
        rows = _coerce_rows(kwargs.get("rows") or kwargs.get("values") or kwargs.get("data"))
        if not rows:
            raise ValueError("insert_table 需要 rows（二维数组）")
        cols = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=cols)
        try:
            table.style = str(kwargs.get("style") or "Table Grid")
        except KeyError:
            pass
        for r, row in enumerate(rows):
            for c in range(cols):
                table.rows[r].cells[c].text = str(row[c]) if c < len(row) and row[c] is not None else ""
        # 默认加粗表头行（除非显式 header=false）。
        if kwargs.get("header", True) and len(rows) > 1:
            for cell in table.rows[0].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
        summary = f"Word 已追加表格：{len(rows)} 行 × {cols} 列"
    elif op in ("insert_image", "add_image", "image", "insert_picture", "图片", "插入图片"):
        image = str(kwargs.get("image") or kwargs.get("path") or kwargs.get("src") or "").strip()
        img_path = Path(image).expanduser()
        if not img_path.is_file():
            raise ValueError(f"图片不存在：{image}")
        width = kwargs.get("width") or kwargs.get("width_inches")
        if width:
            try:
                document.add_picture(str(img_path), width=Inches(float(width)))
            except (TypeError, ValueError):
                document.add_picture(str(img_path))
        else:
            document.add_picture(str(img_path))
        align = kwargs.get("align") or kwargs.get("alignment")
        if align and document.paragraphs:
            _docx_align(document.paragraphs[-1], align)
        summary = f"Word 已插入图片：{img_path.name}"
    elif op in ("add_page_break", "page_break", "分页", "分页符"):
        document.add_page_break()
        summary = "Word 已插入分页符"
    elif op in ("set_font", "font", "默认字体", "字体"):
        from docx.shared import Pt  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        name = str(kwargs.get("font") or kwargs.get("name") or kwargs.get("font_name") or "").strip()
        size = kwargs.get("size") or kwargs.get("font_size")
        style = document.styles["Normal"]
        if name:
            style.font.name = name
            try:
                rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
                for slot in ("w:eastAsia", "w:ascii", "w:hAnsi"):
                    rfonts.set(qn(slot), name)
            except Exception:
                pass
        if size:
            style.font.size = Pt(_to_int(size, 12))
        summary = f"Word 默认字体已设置：{(name + ' ') if name else ''}{size or ''}".strip()
    elif op in ("page_setup", "page", "页面设置", "页边距", "margins"):
        from docx.shared import Inches  # type: ignore
        from docx.enum.section import WD_ORIENT  # type: ignore
        section = document.sections[0]
        for key, attr in (("top", "top_margin"), ("bottom", "bottom_margin"),
                          ("left", "left_margin"), ("right", "right_margin")):
            value = kwargs.get(key)
            if value is None:
                value = kwargs.get(key + "_margin")
            if value is not None:
                setattr(section, attr, Inches(_to_float(value, 1.0)))
        orient = str(kwargs.get("orientation") or "").strip().lower()
        if orient in ("landscape", "横向", "horizontal"):
            section.orientation = WD_ORIENT.LANDSCAPE
            if section.page_width < section.page_height:
                section.page_width, section.page_height = section.page_height, section.page_width
        elif orient in ("portrait", "纵向", "vertical"):
            section.orientation = WD_ORIENT.PORTRAIT
            if section.page_width > section.page_height:
                section.page_width, section.page_height = section.page_height, section.page_width
        summary = "Word 页面设置已更新"
    elif op in ("set_header", "header", "页眉"):
        text = str(kwargs.get("text") or kwargs.get("content") or "")
        section = document.sections[0]
        section.header.is_linked_to_previous = False
        para = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        para.text = text
        align = kwargs.get("align") or kwargs.get("alignment")
        _docx_align(para, align or "center")
        summary = "Word 页眉已设置"
    elif op in ("set_footer", "footer", "页脚"):
        text = str(kwargs.get("text") or kwargs.get("content") or "")
        section = document.sections[0]
        section.footer.is_linked_to_previous = False
        para = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        para.text = text
        align = kwargs.get("align") or kwargs.get("alignment")
        _docx_align(para, align or "center")
        summary = "Word 页脚已设置"
    elif op in ("add_toc", "toc", "目录", "table_of_contents"):
        _docx_add_field(document.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
                        placeholder="（目录：打开或按 F9 更新）")
        _docx_enable_update_fields(document)
        summary = "Word 已插入目录域（打开文档即自动更新页码）"
    elif op in ("add_bookmark", "bookmark", "书签", "标记"):
        name = _docx_safe_bookmark_name(kwargs.get("name") or kwargs.get("bookmark") or kwargs.get("id"))
        if not name:
            raise ValueError("add_bookmark 需要 name（书签名）")
        find = kwargs.get("find") or kwargs.get("target") or kwargs.get("anchor")
        text = kwargs.get("text") or kwargs.get("content")
        if find:
            paragraph = _docx_find_paragraph(document, find)
            if paragraph is None:
                raise ValueError(f"未找到包含「{find}」的段落，无法定位书签")
        elif text is not None:
            paragraph = document.add_paragraph(str(text))
        else:
            raise ValueError("add_bookmark 需要 find（已有段落文本）或 text（新建被标记段落）")
        _docx_wrap_bookmark(paragraph, name, document)
        summary = f"Word 已添加书签：{name}"
    elif op in ("add_cross_reference", "cross_reference", "crossref", "cross_ref",
                "交叉引用", "引用"):
        bookmark = _docx_safe_bookmark_name(kwargs.get("bookmark") or kwargs.get("target")
                                            or kwargs.get("ref") or kwargs.get("name"))
        if not bookmark:
            raise ValueError("add_cross_reference 需要 bookmark（目标书签名）")
        ref_type = str(kwargs.get("ref_type") or kwargs.get("type") or "text").strip().lower()
        paragraph = document.add_paragraph()
        prefix = kwargs.get("prefix") or kwargs.get("before")
        if prefix:
            paragraph.add_run(str(prefix))
        if ref_type in ("page", "页码", "页"):
            _docx_add_field(paragraph, f"PAGEREF {bookmark} \\h", placeholder="?")
        elif ref_type in ("both", "text_page", "全部", "文本和页码"):
            _docx_add_field(paragraph, f"REF {bookmark} \\h", placeholder="?")
            paragraph.add_run("（见第 ")
            _docx_add_field(paragraph, f"PAGEREF {bookmark} \\h", placeholder="?")
            paragraph.add_run(" 页）")
        else:  # 默认引用文本（标题文字 / 题注文字）
            _docx_add_field(paragraph, f"REF {bookmark} \\h", placeholder="?")
        suffix = kwargs.get("suffix") or kwargs.get("after")
        if suffix:
            paragraph.add_run(str(suffix))
        align = kwargs.get("align") or kwargs.get("alignment")
        if align:
            _docx_align(paragraph, align)
        _docx_enable_update_fields(document)
        summary = f"Word 已插入对「{bookmark}」的交叉引用（打开文档即自动更新）"
    elif op in ("add_caption", "caption", "题注", "图注", "表注"):
        label = str(kwargs.get("label") or kwargs.get("category") or "图").strip() or "图"
        text = str(kwargs.get("text") or kwargs.get("content") or kwargs.get("caption") or "").strip()
        separator = str(kwargs.get("separator") or kwargs.get("sep") or " ")
        try:
            paragraph = document.add_paragraph(style="Caption")
        except KeyError:
            paragraph = document.add_paragraph()
        paragraph.add_run(f"{label} ")
        _docx_add_field(paragraph, f"SEQ {label} \\* ARABIC", placeholder="1")
        if text:
            paragraph.add_run(f"{separator}{text}")
        align = kwargs.get("align") or kwargs.get("alignment")
        if align:
            _docx_align(paragraph, align)
        bookmark = kwargs.get("bookmark") or kwargs.get("name")
        bm_note = ""
        if bookmark:
            safe = _docx_safe_bookmark_name(bookmark)
            _docx_wrap_bookmark(paragraph, safe, document)
            bm_note = f"，书签 {safe}"
        _docx_enable_update_fields(document)
        summary = f"Word 已添加{label}题注（自动编号{bm_note}）"
    else:
        raise ValueError(f"Word 不支持的编辑操作：{op}")
    document.save(str(destination))
    return summary


def _coerce_rows(value) -> list[list]:
    """把多种形态的输入归一成二维数组（list[list]）。"""
    if value is None:
        return []
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except (ValueError, TypeError):
            parsed = None
        if isinstance(parsed, list):
            value = parsed
        else:
            rows = []
            for line in value.splitlines():
                line = line.strip()
                if line:
                    rows.append([cell.strip() for cell in re.split(r"[\t,]", line)])
            return rows
    if isinstance(value, (list, tuple)):
        rows = []
        for row in value:
            if isinstance(row, (list, tuple)):
                rows.append(list(row))
            elif isinstance(row, str):
                rows.append([cell.strip() for cell in re.split(r"[\t,]", row)])
            else:
                rows.append([row])
        return rows
    return [[value]]


# ---- Office 编辑：通用小工具（颜色/列号/取整/单元格区域） --------------------
_NAMED_COLORS = {
    "RED": "FFFF0000", "GREEN": "FF00B050", "BLUE": "FF0070C0", "YELLOW": "FFFFFF00",
    "ORANGE": "FFFFA500", "PURPLE": "FF7030A0", "BLACK": "FF000000", "WHITE": "FFFFFFFF",
    "GRAY": "FF808080", "GREY": "FF808080", "DARKBLUE": "FF1F3864", "LIGHTBLUE": "FFDDEBF7",
    "LIGHTGRAY": "FFF2F2F2", "LIGHTGREY": "FFF2F2F2", "PINK": "FFFF66CC", "BROWN": "FF843C0C",
    "红": "FFFF0000", "绿": "FF00B050", "蓝": "FF0070C0", "黄": "FFFFFF00",
    "橙": "FFFFA500", "紫": "FF7030A0", "黑": "FF000000", "白": "FFFFFFFF", "灰": "FF808080",
}


def _norm_hex(value) -> "str | None":
    """把颜色输入归一成 8 位 ARGB（如 FFFF0000）；支持 #RRGGBB / RRGGBB / 颜色名。"""
    if value is None:
        return None
    text = str(value).strip().lstrip("#")
    upper = text.upper()
    if upper in _NAMED_COLORS:
        return _NAMED_COLORS[upper]
    if all(c in "0123456789ABCDEF" for c in upper):
        if len(upper) == 6:
            return "FF" + upper
        if len(upper) == 8:
            return upper
    return None


def _norm_rgb(value) -> "tuple | None":
    """颜色输入 → (r, g, b)，供 python-pptx 的 RGBColor 用。"""
    argb = _norm_hex(value)
    if not argb:
        return None
    rgb = argb[-6:]
    return (int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16))


def _to_int(value, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _to_float(value, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _col_index(value) -> int:
    """列号归一成 1 起的整数：接受 'A' / 'B2' / 3 等形态。"""
    if isinstance(value, (int, float)):
        return max(1, int(value))
    text = str(value or "A").strip().upper()
    letters = "".join(ch for ch in text if ch.isalpha())
    if letters:
        try:
            from openpyxl.utils import column_index_from_string
            return column_index_from_string(letters)
        except Exception:
            return 1
    return max(1, _to_int(text, 1))


def _range_cells(ws, rng: str) -> list:
    """把工作表区域（A1 或 A1:C3）摊平成单元格列表。"""
    selection = ws[rng]
    if not isinstance(selection, tuple):
        return [selection]
    flat = []
    for item in selection:
        if isinstance(item, tuple):
            flat.extend(item)
        else:
            flat.append(item)
    return flat


def _edit_pptx_copy(source: Path, destination: Path, op: str, kwargs: dict) -> str:
    if op in ("add_slide", "new_slide", "append_slide", "新建幻灯片", "添加幻灯片"):
        return _pptx_python_add_slide(source, destination, kwargs)
    if op in ("add_table", "insert_table", "table", "表格", "幻灯片表格"):
        return _pptx_add_table(source, destination, kwargs)
    if op in ("add_image", "add_picture", "insert_image", "insert_picture",
              "image", "picture", "图片", "插入图片"):
        return _pptx_add_image(source, destination, kwargs)
    if op in ("add_textbox", "textbox", "text_box", "文本框", "插入文本"):
        return _pptx_add_textbox(source, destination, kwargs)
    if op in ("add_chart", "chart", "insert_chart", "图表"):
        return _pptx_add_chart(source, destination, kwargs)
    if op in ("delete_slide", "remove_slide", "删除幻灯片"):
        return _pptx_delete_slide(source, destination, kwargs)
    if op in ("set_background", "background", "背景", "设置背景"):
        return _pptx_set_background(source, destination, kwargs)
    if op not in ("replace", "replace_text", "text_replace"):
        raise ValueError(
            "PowerPoint 支持 replace_text / add_slide / add_table / add_image / "
            f"add_textbox / add_chart / delete_slide / set_background，当前操作：{op}")
    find = str(kwargs.get("find") or kwargs.get("search") or "")
    replace = str(kwargs.get("replace") or kwargs.get("text") or "")
    if not find:
        raise ValueError("replace_text 需要 find")
    replacements: dict[str, bytes] = {}
    count = 0
    ns_uri = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with zipfile.ZipFile(source) as zf:
        slide_names = [name for name in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        for name in slide_names:
            root = _xml_text(zf.read(name).decode("utf-8", errors="replace"))
            changed = False
            for node in root.findall(f".//{{{ns_uri}}}t"):
                if node.text and find in node.text:
                    count += node.text.count(find)
                    node.text = node.text.replace(find, replace)
                    changed = True
            if changed:
                replacements[name] = _xml_bytes(root)
    if count == 0:
        raise ValueError(f"未找到文本：{find}")
    _write_zip_copy(source, destination, replacements)
    return f"PowerPoint 文本替换完成：{count} 处"


def _edit_xlsx_copy(source: Path, destination: Path, op: str, kwargs: dict) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise RuntimeError("编辑 Excel 需要 openpyxl") from exc
    shutil.copy2(source, destination)
    wb = openpyxl.load_workbook(destination)
    try:
        sheet_name = str(kwargs.get("sheet") or "").strip()
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        if op in ("set_cell", "cell"):
            cell = str(kwargs.get("cell") or "").strip().upper()
            if not cell:
                raise ValueError("set_cell 需要 cell，例如 A1")
            ws[cell] = kwargs.get("value", "")
            wb.save(destination)
            return f"Excel 单元格已修改：{ws.title}!{cell}"
        if op in ("append_row", "append"):
            values = kwargs.get("values")
            if values is None:
                values = kwargs.get("value", kwargs.get("text", ""))
            if isinstance(values, str):
                values = [v.strip() for v in values.replace("\t", ",").split(",")]
            if not isinstance(values, (list, tuple)):
                values = [values]
            ws.append(list(values))
            wb.save(destination)
            return f"Excel 已追加一行：{ws.title}"
        if op in ("set_formula", "formula", "公式"):
            cell = str(kwargs.get("cell") or "").strip().upper()
            formula = str(kwargs.get("formula") or kwargs.get("value") or "").strip()
            if not cell or not formula:
                raise ValueError("set_formula 需要 cell 和 formula，例如 cell=C1, formula==A1+B1")
            ws[cell] = formula if formula.startswith("=") else f"={formula}"
            wb.save(destination)
            return f"Excel 公式已写入：{ws.title}!{cell}"
        if op in ("write_rows", "insert_table", "set_range", "写入区域"):
            rows = _coerce_rows(kwargs.get("rows") or kwargs.get("values") or kwargs.get("data"))
            if not rows:
                raise ValueError("write_rows 需要 rows（二维数组）")
            start = str(kwargs.get("cell") or kwargs.get("start") or "A1").strip().upper()
            try:
                from openpyxl.utils.cell import coordinate_to_tuple  # type: ignore
                start_row, start_col = coordinate_to_tuple(start)
            except Exception:
                start_row, start_col = 1, 1
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    ws.cell(row=start_row + r, column=start_col + c, value=val)
            wb.save(destination)
            return f"Excel 已写入 {len(rows)} 行（起点 {ws.title}!{start}）"
        if op in ("delete_row", "remove_row", "删除行"):
            try:
                row = int(kwargs.get("row") or kwargs.get("index") or kwargs.get("value"))
            except (TypeError, ValueError) as exc:
                raise ValueError("delete_row 需要 row（行号，从 1 开始）") from exc
            ws.delete_rows(row)
            wb.save(destination)
            return f"Excel 已删除第 {row} 行：{ws.title}"
        if op in ("add_sheet", "new_sheet", "新建工作表"):
            name = str(kwargs.get("name") or kwargs.get("title") or kwargs.get("sheet") or "").strip()
            new_ws = wb.create_sheet(title=name or None)
            wb.save(destination)
            return f"Excel 已新建工作表：{new_ws.title}"
        if op in ("replace", "replace_text", "text_replace"):
            find = str(kwargs.get("find") or kwargs.get("search") or "")
            replace = str(kwargs.get("replace") or kwargs.get("text") or "")
            if not find:
                raise ValueError("replace_text 需要 find")
            count = 0
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and find in cell.value:
                            count += cell.value.count(find)
                            cell.value = cell.value.replace(find, replace)
            if count == 0:
                raise ValueError(f"未找到文本：{find}")
            wb.save(destination)
            return f"Excel 文本替换完成：{count} 处"
        if op in ("insert_row", "insert_rows", "插入行"):
            idx = max(1, _to_int(kwargs.get("row") or kwargs.get("index"), 1))
            count = max(1, _to_int(kwargs.get("count"), 1))
            ws.insert_rows(idx, amount=count)
            values = kwargs.get("values")
            if values is None:
                values = kwargs.get("value")
            if values is not None:
                if isinstance(values, str):
                    values = [v.strip() for v in values.replace("\t", ",").split(",")]
                if not isinstance(values, (list, tuple)):
                    values = [values]
                for c, val in enumerate(values, start=1):
                    ws.cell(row=idx, column=c, value=val)
            wb.save(destination)
            return f"Excel 已在第 {idx} 行插入 {count} 行：{ws.title}"
        if op in ("insert_column", "insert_col", "insert_columns", "插入列"):
            cidx = _col_index(kwargs.get("column") or kwargs.get("col") or kwargs.get("cell") or "A")
            count = max(1, _to_int(kwargs.get("count"), 1))
            ws.insert_cols(cidx, amount=count)
            header = kwargs.get("header") or kwargs.get("name") or kwargs.get("value")
            if header:
                ws.cell(row=1, column=cidx, value=header)
            wb.save(destination)
            return f"Excel 已插入 {count} 列（第 {cidx} 列起）：{ws.title}"
        if op in ("delete_column", "delete_col", "remove_column", "删除列"):
            cidx = _col_index(kwargs.get("column") or kwargs.get("col") or kwargs.get("cell"))
            count = max(1, _to_int(kwargs.get("count"), 1))
            ws.delete_cols(cidx, amount=count)
            wb.save(destination)
            return f"Excel 已删除第 {cidx} 列起 {count} 列：{ws.title}"
        if op in ("format_cell", "format", "style", "style_range", "set_style", "格式", "样式"):
            from openpyxl.styles import Font, PatternFill, Alignment
            rng = str(kwargs.get("range") or kwargs.get("cell") or "").strip().upper()
            if not rng:
                raise ValueError("format_cell 需要 range 或 cell，例如 range=A1:C1")
            cells = _range_cells(ws, rng)
            font_color = _norm_hex(kwargs.get("color") or kwargs.get("font_color"))
            fill_hex = _norm_hex(kwargs.get("fill") or kwargs.get("bg")
                                 or kwargs.get("bg_color") or kwargs.get("background"))
            size = kwargs.get("size") or kwargs.get("font_size")
            bold = kwargs.get("bold")
            italic = kwargs.get("italic")
            halign = kwargs.get("align") or kwargs.get("h_align") or kwargs.get("horizontal")
            valign = kwargs.get("v_align") or kwargs.get("vertical")
            wrap = kwargs.get("wrap")
            num_fmt = kwargs.get("number_format") or kwargs.get("format_code")
            font_touch = any(v is not None for v in (font_color, size, bold, italic))
            align_touch = any(v is not None for v in (halign, valign, wrap))
            for cell in cells:
                if font_touch:
                    base = cell.font
                    cell.font = Font(
                        name=base.name,
                        size=(_to_int(size, 11) if size else base.size),
                        bold=(bool(bold) if bold is not None else base.bold),
                        italic=(bool(italic) if italic is not None else base.italic),
                        color=(font_color if font_color else base.color),
                    )
                if fill_hex:
                    cell.fill = PatternFill(fill_type="solid", fgColor=fill_hex)
                if align_touch:
                    base = cell.alignment
                    cell.alignment = Alignment(
                        horizontal=(str(halign).lower() if halign else base.horizontal),
                        vertical=(str(valign).lower() if valign else base.vertical),
                        wrap_text=(bool(wrap) if wrap is not None else base.wrap_text),
                    )
                if num_fmt:
                    cell.number_format = str(num_fmt)
            wb.save(destination)
            return f"Excel 已设置 {rng} 格式（{len(cells)} 个单元格）：{ws.title}"
        if op in ("merge_cells", "merge", "合并单元格"):
            rng = str(kwargs.get("range") or kwargs.get("cell") or "").strip().upper()
            if ":" not in rng:
                raise ValueError("merge_cells 需要 range，例如 A1:C1")
            ws.merge_cells(rng)
            val = kwargs.get("value") if kwargs.get("value") is not None else kwargs.get("text")
            if val is not None:
                ws[rng.split(":")[0]] = val
            wb.save(destination)
            return f"Excel 已合并单元格：{rng}（{ws.title}）"
        if op in ("unmerge_cells", "unmerge", "取消合并"):
            rng = str(kwargs.get("range") or kwargs.get("cell") or "").strip().upper()
            ws.unmerge_cells(rng)
            wb.save(destination)
            return f"Excel 已取消合并：{rng}"
        if op in ("auto_width", "autofit", "auto_fit", "自动列宽"):
            from openpyxl.utils import get_column_letter
            widths: dict[int, int] = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
            for col_idx, length in widths.items():
                ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(8, length + 2))
            wb.save(destination)
            return f"Excel 已自动调整 {len(widths)} 列列宽：{ws.title}"
        if op in ("set_column_width", "set_width", "column_width", "列宽"):
            from openpyxl.utils import get_column_letter
            letter = get_column_letter(_col_index(kwargs.get("column") or kwargs.get("col") or kwargs.get("cell") or "A"))
            width = _to_float(kwargs.get("width") or kwargs.get("value"), 12)
            ws.column_dimensions[letter].width = width
            wb.save(destination)
            return f"Excel 已设置列 {letter} 宽度 {width}：{ws.title}"
        if op in ("set_row_height", "set_height", "row_height", "行高"):
            row = max(1, _to_int(kwargs.get("row") or kwargs.get("index"), 1))
            height = _to_float(kwargs.get("height") or kwargs.get("value"), 18)
            ws.row_dimensions[row].height = height
            wb.save(destination)
            return f"Excel 已设置第 {row} 行高度 {height}：{ws.title}"
        if op in ("freeze_panes", "freeze", "冻结窗格", "冻结"):
            cell = str(kwargs.get("cell") or kwargs.get("value") or "A2").strip().upper()
            ws.freeze_panes = cell if cell.lower() not in ("none", "off", "") else None
            wb.save(destination)
            return f"Excel 已冻结窗格于 {cell}：{ws.title}"
        if op in ("clear_range", "clear", "清空区域"):
            rng = str(kwargs.get("range") or kwargs.get("cell") or "").strip().upper()
            if not rng:
                raise ValueError("clear_range 需要 range")
            for cell in _range_cells(ws, rng):
                cell.value = None
            wb.save(destination)
            return f"Excel 已清空 {rng}：{ws.title}"
        if op in ("rename_sheet", "rename", "重命名工作表"):
            new = str(kwargs.get("new_name") or kwargs.get("name")
                      or kwargs.get("title") or kwargs.get("value") or "").strip()
            if not new:
                raise ValueError("rename_sheet 需要 new_name")
            old = ws.title
            ws.title = new[:31]
            wb.save(destination)
            return f"Excel 工作表已重命名：{old} → {ws.title}"
        if op in ("delete_sheet", "remove_sheet", "删除工作表"):
            if len(wb.sheetnames) <= 1:
                raise ValueError("不能删除唯一的工作表")
            removed = ws.title
            del wb[removed]
            wb.save(destination)
            return f"Excel 已删除工作表：{removed}"
        if op in ("copy_sheet", "duplicate_sheet", "复制工作表"):
            new_ws = wb.copy_worksheet(ws)
            nm = kwargs.get("new_name") or kwargs.get("name")
            if nm:
                new_ws.title = str(nm)[:31]
            wb.save(destination)
            return f"Excel 已复制工作表：{ws.title} → {new_ws.title}"
        if op in ("add_chart", "chart", "insert_chart", "图表"):
            from openpyxl.chart import BarChart, LineChart, PieChart, Reference
            from openpyxl.utils.cell import range_boundaries
            ctype = str(kwargs.get("chart_type") or kwargs.get("type") or "bar").strip().lower()
            data_ref = str(kwargs.get("data_range") or kwargs.get("data") or kwargs.get("range") or "").strip().upper()
            if ":" not in data_ref:
                raise ValueError("add_chart 需要 data_range（含表头），例如 data_range=A1:B10")
            min_col, min_row, max_col, max_row = range_boundaries(data_ref)
            if "pie" in ctype:
                chart = PieChart()
            elif "line" in ctype:
                chart = LineChart()
            else:
                chart = BarChart()
                chart.type = "bar" if ctype in ("barh", "bar_h", "horizontal") else "col"
            title = kwargs.get("title")
            if title:
                chart.title = str(title)
            data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            cats_ref = str(kwargs.get("categories") or kwargs.get("cats") or "").strip().upper()
            if cats_ref:
                cb = range_boundaries(cats_ref)
                chart.set_categories(Reference(ws, min_col=cb[0], min_row=cb[1], max_col=cb[2], max_row=cb[3]))
            anchor = str(kwargs.get("anchor") or kwargs.get("at") or "H2").strip().upper()
            ws.add_chart(chart, anchor)
            wb.save(destination)
            return f"Excel 已插入 {ctype} 图表（数据 {data_ref}）：{ws.title}"
        raise ValueError(f"Excel 不支持的编辑操作：{op}")
    finally:
        try:
            wb.close()
        except tk.TclError:
            pass


def _xml_bytes(root: ET.Element) -> bytes:
    ET.register_namespace("w", "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
    ET.register_namespace("a", "http://schemas.openxmlformats.org/drawingml/2006/main")
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _write_zip_copy(source: Path, destination: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = replacements.get(info.filename)
            if data is None:
                data = zin.read(info.filename)
            zout.writestr(info, data)


def _xml_text(value: str) -> ET.Element:
    return ET.fromstring(value.encode("utf-8"))


def _extract_docx_text(path: Path) -> str:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    root = _xml_text(xml)
    lines = []
    for para in root.findall(".//w:p", ns):
        text = "".join(node.text or "" for node in para.findall(".//w:t", ns)).strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def _extract_pptx_text(path: Path) -> str:
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    lines = []
    with zipfile.ZipFile(path) as zf:
        slide_names = sorted(
            name for name in zf.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        for i, name in enumerate(slide_names[:40], start=1):
            root = _xml_text(zf.read(name).decode("utf-8", errors="replace"))
            parts = [node.text or "" for node in root.findall(".//a:t", ns)]
            body = " ".join(part.strip() for part in parts if part.strip())
            if body:
                lines.append(f"Slide {i}: {body}")
    return "\n".join(lines)


def _extract_xlsx_text(path: Path) -> str:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets[:8]:
            lines.append(f"[{ws.title}]")
            for row in ws.iter_rows(max_row=80, max_col=12, values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append("\t".join(values).rstrip())
        try:
            wb.close()
        except tk.TclError:
            pass
        return "\n".join(lines)
    except Exception:
        return _extract_xlsx_text_zip(path)


def _extract_xlsx_text_zip(path: Path) -> str:
    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = _xml_text(zf.read("xl/sharedStrings.xml").decode("utf-8", errors="replace"))
            for si in root.findall(".//main:si", ns):
                shared.append("".join(t.text or "" for t in si.findall(".//main:t", ns)))
        sheet_names = sorted(name for name in zf.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
        lines = []
        for i, name in enumerate(sheet_names[:8], start=1):
            lines.append(f"[Sheet{i}]")
            root = _xml_text(zf.read(name).decode("utf-8", errors="replace"))
            for row in root.findall(".//main:row", ns)[:80]:
                values = []
                for cell in row.findall("main:c", ns)[:12]:
                    value = cell.find("main:v", ns)
                    raw = value.text if value is not None else ""
                    if cell.get("t") == "s" and raw.isdigit() and int(raw) < len(shared):
                        raw = shared[int(raw)]
                    values.append(raw or "")
                if any(v.strip() for v in values):
                    lines.append("\t".join(values).rstrip())
        return "\n".join(lines)


def _coerce_bullets(value) -> list[str]:
    """把要点输入归一成字符串列表。"""
    if value is None:
        return []
    if isinstance(value, str):
        return [line.strip() for line in value.splitlines() if line.strip()]
    if isinstance(value, (list, tuple)):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()]


# 图文并茂用图：模型给的 image 可能是本地路径、http(s) 链接或 data URI。统一解析成
# 一个 python-pptx 能直接插入的本地图片文件；下载/解码后的图缓存在此目录复用。
PPTX_IMAGE_MAX_BYTES = 12 * 1024 * 1024  # 单张配图最大 12MB，超限则跳过


def _pptx_image_cache_dir() -> Path:
    # 惰性取目录：OFFICE_OUTPUT_DIR 在本函数之后定义，且会随数据目录切换被重新赋值。
    return OFFICE_OUTPUT_DIR / "_images"


def _normalize_image_bytes(raw: bytes, dest_stem: Path) -> Path | None:
    """把任意图片字节落地成 python-pptx 可用的文件。

    优先用 Pillow 重新编码（可把 WebP 等 pptx 不支持的格式转成 PNG，并顺带校验
    确实是有效图片）；无 Pillow 时退化为按magic字节判别后缀直接写出（不支持的
    格式返回 None）。
    """
    try:
        from PIL import Image  # type: ignore
    except Exception:
        Image = None
    if Image is not None:
        try:
            with Image.open(io.BytesIO(raw)) as img:
                img.load()
                if img.mode in ("P", "RGBA", "LA"):
                    img = img.convert("RGBA")
                    dest = dest_stem.with_suffix(".png")
                    img.save(dest, "PNG")
                else:
                    img = img.convert("RGB")
                    dest = dest_stem.with_suffix(".png")
                    img.save(dest, "PNG")
            return dest
        except Exception:
            return None
    # 无 Pillow：按文件头粗略判别 pptx 支持的格式。
    sig_ext = {
        b"\x89PNG\r\n\x1a\n": ".png",
        b"\xff\xd8\xff": ".jpg",
        b"GIF87a": ".gif",
        b"GIF89a": ".gif",
        b"BM": ".bmp",
    }
    ext = None
    for sig, candidate in sig_ext.items():
        if raw.startswith(sig):
            ext = candidate
            break
    if ext is None:
        return None
    dest = dest_stem.with_suffix(ext)
    try:
        dest.write_bytes(raw)
    except OSError:
        return None
    return dest


def _download_pptx_image(url: str) -> Path | None:
    cache_dir = _pptx_image_cache_dir()
    cache_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha1(url.encode("utf-8", "surrogatepass")).hexdigest()[:16]
    for existing in cache_dir.glob(f"{digest}.*"):
        if existing.is_file() and existing.stat().st_size > 0:
            return existing
    try:
        req = urllib.request.Request(url, headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) PasserAI/1.0",
        })
        with urllib.request.urlopen(req, timeout=20) as resp:
            raw = resp.read(PPTX_IMAGE_MAX_BYTES + 1)
    except Exception:
        return None
    if not raw or len(raw) > PPTX_IMAGE_MAX_BYTES:
        return None
    return _normalize_image_bytes(raw, cache_dir / digest)


def _resolve_pptx_image(image) -> Path | None:
    """把 image 引用解析为本地图片路径：本地文件 / http(s) 链接 / data URI。失败返回 None。"""
    ref = str(image or "").strip()
    if not ref:
        return None
    low = ref.lower()
    if low.startswith("data:image/"):
        try:
            _, b64 = ref.split(",", 1)
            raw = base64.b64decode(b64)
        except Exception:
            return None
        cache_dir = _pptx_image_cache_dir()
        cache_dir.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha1(raw).hexdigest()[:16]
        return _normalize_image_bytes(raw, cache_dir / digest)
    if low.startswith("http://") or low.startswith("https://"):
        return _download_pptx_image(ref)
    path = Path(ref)
    if path.is_file():
        return path
    return None


def _image_aspect_ratio(image_path: Path) -> float | None:
    """返回图片宽/高比；无法读取时返回 None（调用方退化为仅按宽度缩放）。"""
    try:
        from PIL import Image  # type: ignore
    except Exception:
        return None
    try:
        with Image.open(str(image_path)) as img:
            w, h = img.size
        if w > 0 and h > 0:
            return w / h
    except Exception:
        return None
    return None


def _pptx_slide_size_inches(presentation) -> tuple[float, float]:
    try:
        w = (presentation.slide_width or 0) / 914400.0
        h = (presentation.slide_height or 0) / 914400.0
    except Exception:
        w = h = 0.0
    return (w or 10.0, h or 7.5)


def _pptx_place_image_fit(slide, image_path: Path, box_left: float, box_top: float,
                          box_w: float, box_h: float) -> None:
    """把图片按等比缩放居中放进给定矩形框（单位：英寸）。"""
    from pptx.util import Inches
    ratio = _image_aspect_ratio(image_path)
    if ratio:
        if box_w / box_h > ratio:   # 框比图更宽 → 以高为准
            draw_h = box_h
            draw_w = box_h * ratio
        else:                        # 框比图更高 → 以宽为准
            draw_w = box_w
            draw_h = box_w / ratio
        left = box_left + (box_w - draw_w) / 2.0
        top = box_top + (box_h - draw_h) / 2.0
        slide.shapes.add_picture(str(image_path), Inches(left), Inches(top),
                                 Inches(draw_w), Inches(draw_h))
    else:
        # 拿不到宽高比：只给宽度，python-pptx 会按原始比例算高度。
        slide.shapes.add_picture(str(image_path), Inches(box_left), Inches(box_top),
                                 width=Inches(box_w))


def _pptx_add_bullets_box(slide, bullets: list[str], left: float, top: float,
                          width: float, height: float, size: int = 18) -> None:
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(bullets):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = f"• {line}"
        para.font.size = Pt(size)


def _pptx_add_text_image_slide(presentation, title: str, bullets: list[str],
                               image_path: Path, image_left: bool) -> None:
    """图文并茂版式：标题在上，要点与配图左右分栏。image_left 决定图片在左还是右。"""
    slide_w, slide_h = _pptx_slide_size_inches(presentation)
    slide = presentation.slides.add_slide(_pptx_blank_layout(presentation))
    _pptx_add_title(slide, title)
    margin = 0.5
    gap = 0.4
    content_top = 1.5 if title else 0.6
    content_h = max(1.0, slide_h - content_top - 0.5)
    col_w = max(1.0, (slide_w - 2 * margin - gap) / 2.0)
    text_left = margin + col_w + gap if image_left else margin
    image_box_left = margin if image_left else margin + col_w + gap
    _pptx_add_bullets_box(slide, bullets, text_left, content_top, col_w, content_h)
    _pptx_place_image_fit(slide, image_path, image_box_left, content_top, col_w, content_h)


def _pptx_python_add_slide(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    presentation = pptx.Presentation(str(source))
    title = str(kwargs.get("title") or kwargs.get("heading") or "").strip()
    bullets = _coerce_bullets(kwargs.get("bullets") or kwargs.get("content") or kwargs.get("text"))
    # 版式 1 通常是「标题和内容」。
    layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 else presentation.slide_layouts[0]
    slide = presentation.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    if bullets:
        body = None
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx != 0:  # 跳过标题占位符
                body = placeholder
                break
        if body is not None and body.has_text_frame:
            frame = body.text_frame
            frame.text = bullets[0]
            for line in bullets[1:]:
                frame.add_paragraph().text = line
    presentation.save(str(destination))
    return f"PowerPoint 已新增幻灯片：{title or '（无标题）'}（{len(bullets)} 个要点）"


def _pptx_blank_layout(presentation):
    """取一个尽量「空白」的版式：默认模板里 6 号通常是空白版式。"""
    layouts = presentation.slide_layouts
    return layouts[6] if len(layouts) > 6 else layouts[-1]


def _pptx_open_base(template):
    """打开演示文稿基底。

    给了模板（.pptx/.potx 路径）就以模板为底，继承其主题、母版、版式与配色，并清空
    模板自带的幻灯片（只留母版/版式），从而在模板风格上生成全新内容；未给模板则用
    python-pptx 默认空白模板。
    """
    import pptx  # type: ignore
    ref = str(template or "").strip()
    if ref:
        path = Path(ref)
        if path.is_file():
            try:
                presentation = pptx.Presentation(str(path))
                id_list = presentation.slides._sldIdLst  # noqa: SLF001 - 无公开清空 API
                for sld in list(id_list):
                    # 同时断开关系再移除引用：只删 sldId 会把幻灯片部件留在包里，
                    # 保存时产生重名部件警告/损坏文件。
                    try:
                        presentation.part.drop_rel(sld.rId)
                    except Exception:
                        pass
                    id_list.remove(sld)
                return presentation
            except Exception:
                pass
    return pptx.Presentation()


def _pptx_target_slide(presentation, kwargs):
    """按 slide/index 取已存在的幻灯片；未给则新增一张空白幻灯片。返回 (slide, is_new)。"""
    idx = kwargs.get("slide")
    if idx is None:
        idx = kwargs.get("index")
    if idx is not None and str(idx).strip() != "":
        try:
            position = int(idx)
        except (TypeError, ValueError):
            raise ValueError("slide 序号无效（应为整数，从 1 开始）")
        slides = list(presentation.slides)
        if position < 1 or position > len(slides):
            raise ValueError(f"slide 序号超范围（应在 1..{len(slides)}）")
        return slides[position - 1], False
    return presentation.slides.add_slide(_pptx_blank_layout(presentation)), True


def _pptx_add_title(slide, title: str) -> None:
    from pptx.util import Inches, Pt
    if not title:
        return
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.9))
    frame = box.text_frame
    frame.text = title
    frame.paragraphs[0].font.size = Pt(28)
    frame.paragraphs[0].font.bold = True


def _pptx_add_table(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    rows = _coerce_rows(kwargs.get("rows") or kwargs.get("data") or kwargs.get("values"))
    if not rows:
        raise ValueError("add_table 需要 rows（二维数组）")
    presentation = pptx.Presentation(str(source))
    slide = presentation.slides.add_slide(_pptx_blank_layout(presentation))
    title = str(kwargs.get("title") or "").strip()
    _pptx_add_title(slide, title)
    n_rows = len(rows)
    n_cols = max(len(row) for row in rows)
    top = Inches(1.4 if title else 0.6)
    table = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.5), top, Inches(9), Inches(min(6.0, 0.4 * n_rows))).table
    for r, row in enumerate(rows):
        for c in range(n_cols):
            table.cell(r, c).text = "" if c >= len(row) or row[c] is None else str(row[c])
    presentation.save(str(destination))
    return f"PowerPoint 已新增表格幻灯片：{n_rows}×{n_cols}"


def _pptx_add_image(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    image_ref = str(kwargs.get("image") or kwargs.get("path") or kwargs.get("picture") or "").strip()
    image_path = _resolve_pptx_image(image_ref)
    if image_path is None:
        raise ValueError("add_image 需要 image（本地图片路径、http(s) 链接或 data URI，且能成功读取）")
    presentation = pptx.Presentation(str(source))
    slide, _ = _pptx_target_slide(presentation, kwargs)
    left = Inches(_to_float(kwargs.get("left"), 1.0))
    top = Inches(_to_float(kwargs.get("top"), 1.0))
    width = kwargs.get("width")
    height = kwargs.get("height")
    extra = {}
    if width:
        extra["width"] = Inches(_to_float(width, 6.0))
    if height:
        extra["height"] = Inches(_to_float(height, 4.0))
    slide.shapes.add_picture(str(image_path), left, top, **extra)
    presentation.save(str(destination))
    return f"PowerPoint 已插入图片：{image_path.name}"


def _pptx_add_textbox(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    text = str(kwargs.get("text") or kwargs.get("content") or "").strip()
    if not text:
        raise ValueError("add_textbox 需要 text")
    presentation = pptx.Presentation(str(source))
    slide, _ = _pptx_target_slide(presentation, kwargs)
    box = slide.shapes.add_textbox(
        Inches(_to_float(kwargs.get("left"), 1.0)), Inches(_to_float(kwargs.get("top"), 1.0)),
        Inches(_to_float(kwargs.get("width"), 8.0)), Inches(_to_float(kwargs.get("height"), 1.5)))
    frame = box.text_frame
    frame.word_wrap = True
    lines = text.split("\n")
    frame.text = lines[0]
    for line in lines[1:]:
        frame.add_paragraph().text = line
    size = kwargs.get("size") or kwargs.get("font_size")
    color = _norm_rgb(kwargs.get("color") or kwargs.get("font_color"))
    bold = kwargs.get("bold")
    for para in frame.paragraphs:
        if size:
            para.font.size = Pt(_to_int(size, 18))
        if bold is not None:
            para.font.bold = bool(bold)
        if color:
            para.font.color.rgb = RGBColor(*color)
    presentation.save(str(destination))
    return "PowerPoint 已添加文本框"


def _pptx_add_chart(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
        from pptx.util import Inches
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    categories = _coerce_bullets(kwargs.get("categories") or kwargs.get("cats"))
    if not categories:
        raise ValueError("add_chart 需要 categories（分类标签）")

    def _nums(values) -> list:
        if isinstance(values, str):
            values = [v for v in re.split(r"[,\t]", values) if v.strip()]
        return [_to_float(v, 0.0) for v in (values or [])]

    chart_data = CategoryChartData()
    chart_data.categories = categories
    series = kwargs.get("series")
    if isinstance(series, dict):
        for name, values in series.items():
            chart_data.add_series(str(name), _nums(values))
    elif isinstance(series, list):
        for i, item in enumerate(series):
            if isinstance(item, dict):
                chart_data.add_series(str(item.get("name") or f"系列{i + 1}"), _nums(item.get("values")))
            else:
                chart_data.add_series(f"系列{i + 1}", _nums(item))
    else:
        values = kwargs.get("values")
        if values is None:
            raise ValueError("add_chart 需要 series 或 values")
        chart_data.add_series(str(kwargs.get("series_name") or "数据"), _nums(values))

    ctype = str(kwargs.get("chart_type") or kwargs.get("type") or "bar").strip().lower()
    chart_type = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "barh": XL_CHART_TYPE.BAR_CLUSTERED, "horizontal": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE,
    }.get(ctype, XL_CHART_TYPE.COLUMN_CLUSTERED)
    presentation = pptx.Presentation(str(source))
    slide = presentation.slides.add_slide(_pptx_blank_layout(presentation))
    title = str(kwargs.get("title") or "").strip()
    _pptx_add_title(slide, title)
    slide.shapes.add_chart(
        chart_type, Inches(1), Inches(1.5 if title else 1.0), Inches(8), Inches(5), chart_data)
    presentation.save(str(destination))
    return f"PowerPoint 已新增 {ctype} 图表幻灯片"


def _pptx_delete_slide(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    presentation = pptx.Presentation(str(source))
    id_list = presentation.slides._sldIdLst  # noqa: SLF001 - python-pptx 无公开删除 API
    entries = list(id_list)
    position = _to_int(kwargs.get("slide") or kwargs.get("index") or kwargs.get("value"), 0)
    if position < 1 or position > len(entries):
        raise ValueError(f"slide 序号无效（应在 1..{len(entries)}）")
    id_list.remove(entries[position - 1])
    presentation.save(str(destination))
    return f"PowerPoint 已删除第 {position} 张幻灯片"


def _pptx_set_background(source: Path, destination: Path, kwargs: dict) -> str:
    try:
        import pptx  # type: ignore
        from pptx.dml.color import RGBColor
    except ImportError as exc:
        raise RuntimeError("该 PowerPoint 操作需要 python-pptx 库") from exc
    color = _norm_rgb(kwargs.get("color") or kwargs.get("bg") or kwargs.get("value"))
    if not color:
        raise ValueError("set_background 需要 color（如 #1F3864 / red / 蓝）")
    presentation = pptx.Presentation(str(source))
    idx = kwargs.get("slide")
    if idx is None:
        idx = kwargs.get("index")
    if idx is not None and str(idx).strip() != "":
        slides = list(presentation.slides)
        position = _to_int(idx, 0)
        if position < 1 or position > len(slides):
            raise ValueError(f"slide 序号无效（应在 1..{len(slides)}）")
        targets = [slides[position - 1]]
    else:
        targets = list(presentation.slides)
    for slide in targets:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*color)
    presentation.save(str(destination))
    return f"PowerPoint 已设置背景色（{len(targets)} 张幻灯片）"


# ---------------------------------------------------------------------------
# Office 文件：格式转换 + 新建（均输出到 PasserData/OfficeOutputs，绝不覆盖原文件）
# ---------------------------------------------------------------------------
OFFICE_OUTPUT_DIR = AI_DATA_DIR / "OfficeOutputs"


def set_data_dir(new_dir: "str | os.PathLike") -> None:
    """把 ai_chat 的所有数据路径切换到 new_dir（供 Passer 在用户更换数据目录时调用）。"""
    global AI_DATA_DIR, AI_INSTRUCTIONS_FILE, AI_MEMORY_FILE, AI_CONVERSATIONS_FILE
    global AI_OPERATIONS_FILE, AI_USAGE_FILE, AI_SKILLS_DIR, AI_TOOLS_DIR, OFFICE_OUTPUT_DIR
    global _SKILL_CACHE_SIGNATURE, _SKILL_CACHE
    AI_DATA_DIR = Path(new_dir)
    AI_INSTRUCTIONS_FILE = AI_DATA_DIR / "AI_INSTRUCTIONS.md"
    AI_MEMORY_FILE = AI_DATA_DIR / "ai_memory.json"
    AI_CONVERSATIONS_FILE = AI_DATA_DIR / "ai_conversations.json"
    AI_OPERATIONS_FILE = AI_DATA_DIR / "ai_operations.json"
    AI_USAGE_FILE = AI_DATA_DIR / "ai_usage.json"
    AI_SKILLS_DIR = AI_DATA_DIR / "AISkills"
    AI_TOOLS_DIR = AI_DATA_DIR / "AITools"
    OFFICE_OUTPUT_DIR = AI_DATA_DIR / "OfficeOutputs"
    # 技能目录变了，作废缓存签名以便重新扫描。
    _SKILL_CACHE_SIGNATURE = ()
    _SKILL_CACHE = []


def _office_output_dest(stem: str, suffix: str) -> Path:
    OFFICE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    safe_stem = re.sub(r"[^\w一-鿿.-]+", "_", str(stem or "")).strip("._") or "office"
    if not suffix.startswith("."):
        suffix = "." + suffix
    candidate = OFFICE_OUTPUT_DIR / f"{safe_stem}{suffix}"
    index = 2
    while candidate.exists():
        candidate = OFFICE_OUTPUT_DIR / f"{safe_stem}_{index}{suffix}"
        index += 1
    return candidate


def _find_soffice() -> Path | None:
    """查找 LibreOffice 的 soffice 启动器（PDF/图片转换需要）。"""
    for name in ("soffice.exe", "soffice.com", "soffice"):
        found = shutil.which(name)
        if found:
            return Path(found)
    roots = [os.environ.get(key) for key in ("ProgramW6432", "ProgramFiles", "ProgramFiles(x86)", "LOCALAPPDATA")]
    for value in roots:
        if not value:
            continue
        root = Path(value)
        if not root.exists():
            continue
        for pattern in ("LibreOffice/program/soffice.exe", "LibreOffice */program/soffice.exe"):
            try:
                matches = sorted(root.glob(pattern), key=lambda p: len(str(p)))
            except OSError:
                matches = []
            for candidate in matches:
                if candidate.is_file():
                    return candidate
    return None


def _soffice_convert(source: Path, target_filter: str, suffix: str) -> Path:
    """用 LibreOffice headless 转换文档，返回输出文件路径。"""
    soffice = _find_soffice()
    if soffice is None:
        raise RuntimeError("转换为 PDF/图片需要安装 LibreOffice（未检测到 soffice）。")
    OFFICE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    flags = 0
    if os.name == "nt":
        flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    cmd = [str(soffice), "--headless", "--convert-to", target_filter,
           "--outdir", str(OFFICE_OUTPUT_DIR), str(source)]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180, creationflags=flags)
    produced = OFFICE_OUTPUT_DIR / (source.stem + suffix)
    if not produced.is_file():
        raise RuntimeError(f"LibreOffice 转换失败：{proc.stderr.strip() or proc.stdout.strip() or '无输出'}")
    final = _office_output_dest(source.stem, suffix)
    if produced != final:
        shutil.move(str(produced), str(final))
    return final


def _pdf_to_images(pdf_path: Path, stem: str, fmt: str = "png", scale: float = 2.0) -> list[Path]:
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PDF 转图片需要 pypdfium2 库") from exc
    outputs: list[Path] = []
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=scale)
            image = bitmap.to_pil()
            if fmt in ("jpg", "jpeg") and image.mode in ("RGBA", "P", "LA"):
                image = image.convert("RGB")
            dest = _office_output_dest(f"{stem}_p{index + 1}", "." + fmt)
            image.save(dest)
            outputs.append(dest)
    finally:
        pdf.close()
    return outputs


def office_convert(target: str, to: str) -> list[Path]:
    """把本地 Office/PDF 文件转换为另一种格式，输出到 OfficeOutputs，返回输出路径列表。"""
    source = Path(str(target or "")).expanduser()
    if not source.is_file():
        raise ValueError(f"文件不存在：{source}")
    suffix = source.suffix.lower()
    fmt = str(to or "").strip().lower().lstrip(".")
    aliases = {"jpeg": "jpg", "image": "png", "images": "png", "picture": "png",
               "markdown": "md", "text": "txt", "excel": "csv", "spreadsheet": "csv"}
    fmt = aliases.get(fmt, fmt)
    if not fmt:
        raise ValueError("缺少目标格式 to（如 pdf/csv/txt/png）")

    # CSV：仅电子表格，每个工作表导出一个 CSV。
    if fmt == "csv":
        if suffix not in (".xlsx", ".xlsm"):
            raise ValueError("仅 Excel(.xlsx/.xlsm) 可转 CSV")
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise RuntimeError("转换 Excel 需要 openpyxl") from exc
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        outputs: list[Path] = []
        try:
            for ws in wb.worksheets:
                dest = _office_output_dest(f"{source.stem}_{ws.title}", ".csv")
                with open(dest, "w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.writer(handle)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow(["" if v is None else v for v in row])
                outputs.append(dest)
        finally:
            try:
                wb.close()
            except Exception:
                pass
        if not outputs:
            raise RuntimeError("未导出任何工作表")
        return outputs

    # 纯文本 / Markdown：基于本地文本提取。
    if fmt in ("txt", "md"):
        if suffix == ".docx":
            text = _extract_docx_text(source)
        elif suffix in (".xlsx", ".xlsm"):
            text = _extract_xlsx_text(source)
        elif suffix in (".pptx", ".pptm"):
            text = _extract_pptx_text(source)
        elif suffix == ".pdf":
            text = _pdf_extract_text(source)
        else:
            raise ValueError(f"暂不支持把 {suffix} 转为 {fmt}")
        dest = _office_output_dest(source.stem, "." + fmt)
        dest.write_text(text or "", encoding="utf-8")
        return [dest]

    # PDF：Office 文档经 LibreOffice 转换。
    if fmt == "pdf":
        if suffix == ".pdf":
            raise ValueError("文件已是 PDF")
        if suffix not in OFFICE_ATTACHMENT_EXTS:
            raise ValueError(f"暂不支持把 {suffix} 转为 PDF")
        return [_soffice_convert(source, "pdf", ".pdf")]

    # 图片：PDF 直接渲染；Office 先转 PDF 再渲染。
    if fmt in ("png", "jpg"):
        if suffix == ".pdf":
            return _pdf_to_images(source, source.stem, fmt)
        if suffix in OFFICE_ATTACHMENT_EXTS:
            pdf_path = _soffice_convert(source, "pdf", ".pdf")
            return _pdf_to_images(pdf_path, source.stem, fmt)
        raise ValueError(f"暂不支持把 {suffix} 转为图片")

    raise ValueError(f"不支持的目标格式：{fmt}（可用 pdf/csv/txt/md/png/jpg）")


def _pdf_extract_text(path: Path) -> str:
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError as exc:
        raise RuntimeError("读取 PDF 文本需要 pypdfium2") from exc
    pdf = pdfium.PdfDocument(str(path))
    parts: list[str] = []
    try:
        for index in range(len(pdf)):
            textpage = pdf[index].get_textpage()
            parts.append(textpage.get_text_range())
    finally:
        pdf.close()
    return "\n".join(parts)


def office_create(fmt: str, name: str, spec: dict) -> Path:
    """从零生成 Office 文件（docx/xlsx/pptx），返回输出路径。"""
    kind = str(fmt or "").strip().lower().lstrip(".")
    aliases = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx",
               "spreadsheet": "xlsx", "powerpoint": "pptx", "ppt": "pptx",
               "presentation": "pptx", "slides": "pptx"}
    kind = aliases.get(kind, kind)
    stem = str(name or "").strip() or "新建文档_" + datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = os.path.splitext(stem)[0] or stem
    if kind == "docx":
        return _create_docx(stem, spec)
    if kind == "xlsx":
        return _create_xlsx(stem, spec)
    if kind == "pptx":
        return _create_pptx(stem, spec)
    raise ValueError(f"不支持新建该类型：{fmt}（可用 docx/xlsx/pptx）")


def _docx_open_base(template):
    """打开 Word 文档基底。

    给了模板（.docx/.dotx 路径）就以模板为底，继承其样式、主题字体、页面设置与页眉
    页脚，并清空模板正文（仅保留节属性 sectPr，即页边距/页眉页脚引用），从而在模板
    风格上写入全新内容；未给模板则用 python-docx 默认空白文档。
    """
    import docx  # type: ignore
    from docx.oxml.ns import qn
    ref = str(template or "").strip()
    if ref:
        path = Path(ref)
        if path.is_file():
            try:
                document = docx.Document(str(path))
                body = document.element.body
                for child in list(body):
                    if child.tag == qn("w:sectPr"):
                        continue
                    body.remove(child)
                return document
            except Exception:
                pass
    return docx.Document()


def _create_docx(stem: str, spec: dict) -> Path:
    try:
        import docx  # type: ignore  # noqa: F401 - 触发缺库时友好报错
    except ImportError as exc:
        raise RuntimeError("新建 Word 需要 python-docx 库") from exc
    document = _docx_open_base(spec.get("template") or spec.get("base"))
    title = str(spec.get("title") or spec.get("heading") or "").strip()
    if title:
        try:
            document.add_heading(title, level=0)
        except (KeyError, ValueError):
            # 模板可能缺少内置「Title」样式：退化为加粗段落，避免整篇生成失败。
            run = document.add_paragraph().add_run(title)
            run.bold = True
    paragraphs = spec.get("paragraphs")
    if paragraphs is None:
        content = spec.get("content") or spec.get("text") or ""
        paragraphs = [block.strip() for block in str(content).split("\n\n") if block.strip()] \
            if isinstance(content, str) else _coerce_bullets(content)
    elif isinstance(paragraphs, str):
        paragraphs = [block.strip() for block in paragraphs.split("\n\n") if block.strip()]
    for para in paragraphs or []:
        document.add_paragraph(str(para))
    rows = _coerce_rows(spec.get("rows") or spec.get("table"))
    if rows:
        cols = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=cols)
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        for r, row in enumerate(rows):
            for c in range(cols):
                table.rows[r].cells[c].text = str(row[c]) if c < len(row) else ""
    dest = _office_output_dest(stem, ".docx")
    document.save(str(dest))
    return dest


def _excel_coerce_value(value):
    """把单元格输入归一：以「=」开头当作公式原样写入（openpyxl 会按公式处理），
    纯数字串转成数值（便于套用千分位/小数格式），其余按文本。保留前导零的串当文本。"""
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    s = str(value).strip()
    if not s:
        return None
    if s.startswith("="):
        return s
    if re.fullmatch(r"-?\d+", s) and not (len(s.lstrip("-")) > 1 and s.lstrip("-").startswith("0")):
        try:
            return int(s)
        except ValueError:
            pass
    if re.fullmatch(r"-?\d*\.\d+", s):
        try:
            return float(s)
        except ValueError:
            pass
    return s


def _xlsx_open_base(template, sheet_name: str):
    """取工作簿与目标工作表：给了模板就以模板为底（继承主题/命名样式），清空目标表后
    重新写入，从而「按模板生成相似的新表」；未给模板则新建空白工作簿。"""
    import openpyxl  # type: ignore
    ref = str(template or "").strip()
    wb = None
    if ref:
        path = Path(ref)
        if path.is_file():
            try:
                wb = openpyxl.load_workbook(str(path))
            except Exception:
                wb = None
    if wb is None:
        wb = openpyxl.Workbook()
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        if sheet_name:
            ws.title = sheet_name[:31]
    # 清空目标表已有内容，只借用模板的主题/样式，避免与旧数据串台。仅当表里确有单元格
    # 时才删除（对空表调用 delete_rows 反而会扰乱内部状态），并把 append 指针复位到行首，
    # 规避 openpyxl 删除后不重置 _current_row 导致首行后凭空多出一行的问题。
    if getattr(ws, "_cells", None):
        ws.delete_rows(1, ws.max_row)
    try:
        ws._current_row = 0  # noqa: SLF001 - 复位 append 指针
    except Exception:
        pass
    return wb, ws


def _xlsx_add_chart(ws, chart_spec: dict, header_row: int, first_data_row: int,
                    last_data_row: int, n_cols: int) -> None:
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    ctype = str(chart_spec.get("type") or chart_spec.get("chart_type") or "bar").strip().lower()
    if ctype in ("pie", "饼", "饼图"):
        chart = PieChart()
    elif ctype in ("line", "折线", "折线图"):
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if ctype in ("barh", "horizontal", "条形", "条形图") else "col"
    chart.title = str(chart_spec.get("title") or "").strip() or None
    if last_data_row < first_data_row or n_cols < 2:
        return  # 数据不足以画图
    # 约定：第 1 列为分类标签，第 2..N 列为数值系列（含表头作系列名）。
    data = Reference(ws, min_col=2, max_col=n_cols, min_row=header_row, max_row=last_data_row)
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 16
    from openpyxl.utils import get_column_letter
    anchor = f"{get_column_letter(n_cols + 2)}{header_row}"
    ws.add_chart(chart, anchor)


def _create_xlsx(stem: str, spec: dict) -> Path:
    try:
        import openpyxl  # type: ignore  # noqa: F401 - 触发缺库时友好报错
    except ImportError as exc:
        raise RuntimeError("新建 Excel 需要 openpyxl 库") from exc
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    sheet_name = str(spec.get("sheet") or "").strip()
    wb, ws = _xlsx_open_base(spec.get("template") or spec.get("base"), sheet_name)

    headers = spec.get("headers")
    if isinstance(headers, str):
        headers = [h.strip() for h in re.split(r"[\t,]", headers)]
    headers = [str(h) for h in headers] if headers else []

    rows = _coerce_rows(spec.get("rows") or spec.get("data") or spec.get("values"))
    n_cols = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
    if n_cols == 0:
        raise ValueError("新建 Excel 需要 headers 或 rows 参数")

    thin = Side(style="thin", color="FFD0D7E5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="FF1F3864")
    zebra_fill = PatternFill(fill_type="solid", fgColor="FFF2F5FB")

    header_row = 0
    if headers:
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        # 直接拼坐标，切勿用 ws.cell(row=..) 取坐标——访问单元格会推进 append 指针，
        # 导致表头与数据之间凭空多出一行空行。
        ws.freeze_panes = f"A{header_row + 1}"

    first_data_row = None
    numeric_cols: set[int] = set()
    for row in rows:
        values = [_excel_coerce_value(row[c] if c < len(row) else None) for c in range(n_cols)]
        ws.append(values)
        r = ws.max_row
        if first_data_row is None:
            first_data_row = r
        zebra = ((r - first_data_row) % 2 == 1)
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if zebra:
                cell.fill = zebra_fill
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.##"
                numeric_cols.add(c)
    if first_data_row is None:
        first_data_row = (header_row or 0) + 1
    last_data_row = ws.max_row

    # 合计行：对数值列写 SUM 公式（演示「公式生成」能力）。
    if rows and (spec.get("total_row") or spec.get("totals") or spec.get("total")) and numeric_cols:
        r = last_data_row + 1
        total_label = str(spec.get("total_label") or "合计")
        ws.cell(row=r, column=1, value=total_label).font = Font(bold=True)
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if c in numeric_cols:
                col = get_column_letter(c)
                cell.value = f"=SUM({col}{first_data_row}:{col}{last_data_row})"
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.##"

    # 自动列宽，避免内容被截断。
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col_idx, length in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(8, length + 2))

    chart_spec = spec.get("chart")
    if isinstance(chart_spec, dict) and headers and rows:
        try:
            _xlsx_add_chart(ws, chart_spec, header_row or first_data_row,
                            first_data_row, last_data_row, n_cols)
        except Exception:
            pass  # 图表为增强项，失败不影响数据表生成

    dest = _office_output_dest(stem, ".xlsx")
    wb.save(str(dest))
    return dest


def _create_pptx(stem: str, spec: dict) -> Path:
    try:
        import pptx  # type: ignore  # noqa: F401 - 仅用于触发缺库时的友好报错
    except ImportError as exc:
        raise RuntimeError("新建 PowerPoint 需要 python-pptx 库") from exc
    presentation = _pptx_open_base(spec.get("template") or spec.get("base"))
    slides = spec.get("slides")
    title = str(spec.get("title") or "").strip()
    subtitle = str(spec.get("subtitle") or "").strip()
    if title:
        layout = presentation.slide_layouts[0]  # 标题幻灯片
        slide = presentation.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        if subtitle and len(slide.placeholders) > 1:
            slide.placeholders[1].text = subtitle
    if isinstance(slides, str):
        # 大纲文本：以 # 开头的行为标题，其余为要点。
        parsed: list[dict] = []
        for line in slides.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("#"):
                parsed.append({"title": stripped.lstrip("# ").strip(), "bullets": []})
            elif parsed:
                parsed[-1]["bullets"].append(stripped)
            else:
                parsed.append({"title": stripped, "bullets": []})
        slides = parsed
    for entry in slides or []:
        if not isinstance(entry, dict):
            entry = {"title": str(entry), "bullets": []}
        title_text = str(entry.get("title") or "").strip()
        bullets = _coerce_bullets(entry.get("bullets") or entry.get("content") or entry.get("text"))
        image_path = _resolve_pptx_image(entry.get("image") or entry.get("picture"))
        # 图片在左还是右：image_position/layout 给 "left" 时图在左，默认在右。
        position = str(entry.get("image_position") or entry.get("layout") or "").strip().lower()
        image_left = position in ("left", "图左", "左", "image_left")

        if image_path and bullets:
            # 图文并茂：标题在上，要点与配图左右分栏。
            _pptx_add_text_image_slide(presentation, title_text, bullets, image_path, image_left)
            continue
        if image_path:
            # 纯图片页：空白版式 + 标题 + 居中大图（等比适配）。
            slide = presentation.slides.add_slide(_pptx_blank_layout(presentation))
            _pptx_add_title(slide, title_text)
            slide_w, slide_h = _pptx_slide_size_inches(presentation)
            top = 1.5 if title_text else 0.6
            _pptx_place_image_fit(slide, image_path, 0.7, top,
                                  slide_w - 1.4, max(1.0, slide_h - top - 0.5))
            continue
        layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 \
            else presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title_text
        if bullets:
            body = None
            for placeholder in slide.placeholders:
                if placeholder.placeholder_format.idx != 0:
                    body = placeholder
                    break
            if body is not None and body.has_text_frame:
                frame = body.text_frame
                frame.text = bullets[0]
                for line in bullets[1:]:
                    frame.add_paragraph().text = line
    if not presentation.slides:
        raise ValueError("新建 PPT 需要 title 或 slides 参数")
    dest = _office_output_dest(stem, ".pptx")
    presentation.save(str(dest))
    return dest
