from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import ctypes
import threading
import time
import tkinter as tk
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tkinter import messagebox
from typing import Callable

from clicker_tool import ClickerTheme


AIRA_TITLE = "Aira"
HISTORY_LIMIT = 200
WECHAT_APP_TOKENS = ("wechat", "weixin", "tencent.mm", "微信")
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
CLASS_GROUP_TERMS = ("班", "班级", "班群", "同学", "辅导员", "班主任", "class")
EXCEL_REQUEST_TERMS = ("excel", "xlsx", "表格", "登记表", "统计表", "电子表格")
FILL_REQUEST_TERMS = ("填写", "填报", "登记", "统计", "收集", "补充", "上报")
AIRA_NOTIFY_MODE_LABELS = {
    "passer": "Passer 内通知",
    "windows": "Windows 通知",
}
AIRA_NOTIFY_MODE_BY_LABEL = {label: key for key, label in AIRA_NOTIFY_MODE_LABELS.items()}
USAGE_POLL_SECONDS = 30.0
USAGE_IDLE_RESET_SECONDS = 5 * 60.0
USAGE_REMIND_AFTER_SECONDS = 60 * 60.0
USAGE_SAVE_INTERVAL_SECONDS = 5 * 60.0


def normalize_aira_notify_mode(value) -> str:
    raw = str(value or "").strip()
    if raw in AIRA_NOTIFY_MODE_LABELS:
        return raw
    if raw in AIRA_NOTIFY_MODE_BY_LABEL:
        return AIRA_NOTIFY_MODE_BY_LABEL[raw]
    if raw.casefold() in {"passer", "inside", "internal", "内置", "内部", "passer内通知"}:
        return "passer"
    return "windows"


def windows_idle_seconds() -> float:
    """Return seconds since the last keyboard or mouse input on Windows."""
    if os.name != "nt":
        return 0.0
    import ctypes

    class LastInputInfo(ctypes.Structure):
        _fields_ = (("cbSize", ctypes.c_uint), ("dwTime", ctypes.c_uint))

    info = LastInputInfo()
    info.cbSize = ctypes.sizeof(info)
    if not ctypes.windll.user32.GetLastInputInfo(ctypes.byref(info)):
        raise OSError("无法读取 Windows 最后输入时间。")
    tick = int(ctypes.windll.kernel32.GetTickCount())
    return float((tick - int(info.dwTime)) & 0xFFFFFFFF) / 1000.0


def _format_usage_duration(seconds: float) -> str:
    minutes = max(0, int(float(seconds) // 60))
    hours, minutes = divmod(minutes, 60)
    if hours and minutes:
        return f"{hours} 小时 {minutes} 分钟"
    if hours:
        return f"{hours} 小时"
    return f"{minutes} 分钟"


@dataclass(frozen=True)
class WeChatNotification:
    notification_id: int
    app_id: str
    sender: str
    text: str
    created_at: str
    fingerprint: str


def _notification_text(notifications: list[WeChatNotification]) -> str:
    return "\n".join(
        f"{item.sender}\n{item.text}" for item in notifications
        if item.sender or item.text
    )


def looks_like_class_excel_request(notifications: list[WeChatNotification]) -> bool:
    text = _notification_text(notifications).casefold()
    group_channel = any(item.app_id == "wechat-window-group" for item in notifications)
    return (
        (group_channel or any(term.casefold() in text for term in CLASS_GROUP_TERMS))
        and any(term.casefold() in text for term in EXCEL_REQUEST_TERMS)
        and any(term.casefold() in text for term in FILL_REQUEST_TERMS)
    )


def _explicit_excel_hint(text: str) -> str:
    matches = re.findall(r"[^\\/:*?\"<>|\r\n]{1,90}\.(?:xlsx|xlsm)", text, flags=re.IGNORECASE)
    return Path(matches[-1].strip()).name if matches else ""


def parse_aira_model_result(text: str, notifications: list[WeChatNotification]) -> dict:
    """Normalize Aira's JSON response and locally gate Excel-fill suggestions."""
    raw = str(text or "").strip()
    candidate = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    payload: dict = {}
    try:
        payload = json.loads(candidate)
    except (TypeError, ValueError, json.JSONDecodeError):
        start, end = candidate.find("{"), candidate.rfind("}")
        if start >= 0 and end > start:
            try:
                payload = json.loads(candidate[start:end + 1])
            except (TypeError, ValueError, json.JSONDecodeError):
                payload = {}

    summary = str(payload.get("summary") or raw).strip()[:6000]
    result = {"summary": summary or "收到微信新消息。", "excel_preview": None}
    if not looks_like_class_excel_request(notifications):
        return result

    source = payload.get("excel_fill_preview") or payload.get("excel_preview")
    if not isinstance(source, dict):
        source = {}
    fields: list[dict] = []
    raw_fields = source.get("fields") or []
    if isinstance(raw_fields, dict):
        raw_fields = [{"name": key, "value": value} for key, value in raw_fields.items()]
    if isinstance(raw_fields, list):
        for item in raw_fields[:12]:
            if not isinstance(item, dict):
                continue
            name = re.sub(r"\s+", " ", str(item.get("name") or item.get("field") or "")).strip()[:80]
            if not name:
                continue
            fields.append({
                "name": name,
                "value": str(item.get("value") or "").strip()[:500],
                "required": bool(item.get("required", True)),
            })
    full_text = _notification_text(notifications)
    result["excel_preview"] = {
        "title": str(source.get("title") or "班群 Excel 填写预览").strip()[:100],
        "file_hint": str(source.get("file_hint") or _explicit_excel_hint(full_text)).strip()[:180],
        "sheet": str(source.get("sheet") or "").strip()[:80],
        "fields": fields,
        "reason": str(source.get("reason") or "检测到班群正在收集 Excel 表格内容。" ).strip()[:240],
    }
    return result


def _normalized_header(value) -> str:
    return re.sub(r"[\s:：()（）\[\]【】_-]+", "", str(value or "")).casefold()


def _match_headers(ws, fields: list[dict]) -> tuple[int, dict[str, int], int]:
    field_names = [_normalized_header(item.get("name")) for item in fields]
    best_row, best_map, best_score = 0, {}, 0
    max_row = min(max(ws.max_row, 1), 30)
    max_col = min(max(ws.max_column, 1), 100)
    for row in range(1, max_row + 1):
        headers = {
            _normalized_header(ws.cell(row=row, column=column).value): column
            for column in range(1, max_col + 1)
            if _normalized_header(ws.cell(row=row, column=column).value)
        }
        mapping: dict[str, int] = {}
        for original, normalized in zip(fields, field_names):
            if not normalized:
                continue
            column = headers.get(normalized)
            if column is None:
                candidates = [
                    (header, candidate_column) for header, candidate_column in headers.items()
                    if normalized in header or header in normalized
                ]
                if candidates:
                    candidates.sort(key=lambda pair: abs(len(pair[0]) - len(normalized)))
                    column = candidates[0][1]
            if column is not None:
                mapping[str(original.get("name"))] = column
        if len(mapping) > best_score:
            best_row, best_map, best_score = row, mapping, len(mapping)
    return best_row, best_map, best_score


def apply_excel_fill_preview(source: str | Path, preview: dict, output_dir: str | Path) -> dict:
    """Fill a matched blank row in a copied workbook and return the output details."""
    source_path = Path(source).expanduser().resolve()
    if not source_path.is_file() or source_path.suffix.lower() not in EXCEL_EXTENSIONS:
        raise ValueError("请选择 .xlsx 或 .xlsm 工作簿。")
    fields = [dict(item) for item in preview.get("fields") or [] if isinstance(item, dict)]
    fields = [item for item in fields if str(item.get("name") or "").strip()]
    if not fields:
        raise ValueError("填写预览中没有可写入字段。")
    if not any(str(item.get("value") or "").strip() for item in fields):
        raise ValueError("请先在填写预览中补充至少一个值。")
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise RuntimeError("自动填写 Excel 需要 openpyxl。") from exc

    destination_dir = Path(output_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_stem = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", source_path.stem).strip("._") or "Excel"
    destination = destination_dir / f"{safe_stem}_Aira填写_{stamp}{source_path.suffix}"
    index = 2
    while destination.exists():
        destination = destination_dir / f"{safe_stem}_Aira填写_{stamp}_{index}{source_path.suffix}"
        index += 1
    shutil.copy2(source_path, destination)

    workbook = openpyxl.load_workbook(destination, keep_vba=source_path.suffix.lower() == ".xlsm")
    try:
        preferred = str(preview.get("sheet") or "").strip()
        sheets = [workbook[preferred]] if preferred in workbook.sheetnames else list(workbook.worksheets)
        ranked = []
        for worksheet in sheets:
            header_row, mapping, score = _match_headers(worksheet, fields)
            ranked.append((score, worksheet, header_row, mapping))
        score, worksheet, header_row, mapping = max(ranked, key=lambda item: item[0])
        if score <= 0:
            raise ValueError("未在工作簿前 30 行找到与预览字段匹配的表头。")

        matched_columns = list(mapping.values())
        target_row = max(header_row + 1, 2)
        search_limit = max(worksheet.max_row + 2, target_row + 200)
        while target_row <= search_limit and any(
            worksheet.cell(row=target_row, column=column).value not in (None, "")
            for column in matched_columns
        ):
            target_row += 1
        if target_row > search_limit:
            target_row = worksheet.max_row + 1

        written, unmatched = [], []
        for field in fields:
            name = str(field.get("name") or "").strip()
            column = mapping.get(name)
            if column is None:
                unmatched.append(name)
                continue
            value = field.get("value", "")
            if value in (None, ""):
                continue
            worksheet.cell(row=target_row, column=column, value=str(value))
            written.append(name)
        if not written:
            raise ValueError("预览值没有匹配到可写入的 Excel 表头。")
        workbook.save(destination)
        return {
            "path": str(destination),
            "sheet": worksheet.title,
            "row": target_row,
            "written": written,
            "unmatched": unmatched,
        }
    except Exception:
        try:
            destination.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    finally:
        workbook.close()


class WeChatWindowNotificationReader:
    """Read unread conversation previews from the visible desktop Weixin window."""

    PROCESS_NAMES = {"weixin.exe", "wechat.exe", "wechatappex.exe"}

    def __init__(self, app) -> None:
        self.app = app
        self._running = False
        self.last_window_handle = 0
        self.last_capture_at = ""
        self.last_unread_count = 0

    @staticmethod
    def _wechat_process_ids() -> set[int]:
        try:
            import psutil  # type: ignore
        except ImportError:
            return set()
        result: set[int] = set()
        for process in psutil.process_iter(["pid", "name"]):
            try:
                name = str(process.info.get("name") or "").casefold()
                if name in WeChatWindowNotificationReader.PROCESS_NAMES:
                    result.add(int(process.info["pid"]))
            except (OSError, ValueError, TypeError):
                continue
        return result

    @classmethod
    def _find_window(cls) -> int:
        try:
            import win32gui  # type: ignore
            import win32process  # type: ignore
        except ImportError:
            return 0
        pids = cls._wechat_process_ids()
        if not pids:
            return 0
        candidates: list[tuple[int, int]] = []

        def collect(hwnd, _extra) -> bool:
            try:
                _thread_id, process_id = win32process.GetWindowThreadProcessId(hwnd)
                if process_id not in pids or not win32gui.IsWindowVisible(hwnd):
                    return True
                left, top, right, bottom = win32gui.GetWindowRect(hwnd)
                width, height = right - left, bottom - top
                if width < 360 or height < 280:
                    return True
                candidates.append((width * height, int(hwnd)))
            except Exception:
                pass
            return True

        try:
            win32gui.EnumWindows(collect, None)
        except Exception:
            return 0
        return max(candidates, default=(0, 0))[1]

    def access_status(self) -> str:
        if not self._wechat_process_ids():
            return "wechat_not_running"
        return "allowed" if self._find_window() else "window_hidden"

    def ensure_access(self) -> str:
        return self.access_status()

    def start_channel(self) -> None:
        status = self.ensure_access()
        if status == "wechat_not_running":
            raise RuntimeError("未检测到正在运行的桌面微信。")
        if status != "allowed":
            raise RuntimeError("请先打开微信主界面并保持窗口未最小化。")
        self._running = True

    def stop_channel(self) -> None:
        self._running = False

    @staticmethod
    def _capture_window(hwnd: int):
        try:
            import win32gui  # type: ignore
            import win32ui  # type: ignore
            from ctypes import windll
            from PIL import Image, ImageGrab
        except ImportError as exc:
            raise RuntimeError("微信窗口读取需要 pywin32 和 Pillow。") from exc

        left, top, right, bottom = win32gui.GetWindowRect(hwnd)
        width, height = right - left, bottom - top
        if width <= 0 or height <= 0:
            raise RuntimeError("微信窗口尺寸无效。")

        image = None
        window_dc = save_dc = bitmap = None
        hwnd_dc = None
        try:
            hwnd_dc = win32gui.GetWindowDC(hwnd)
            window_dc = win32ui.CreateDCFromHandle(hwnd_dc)
            save_dc = window_dc.CreateCompatibleDC()
            bitmap = win32ui.CreateBitmap()
            bitmap.CreateCompatibleBitmap(window_dc, width, height)
            save_dc.SelectObject(bitmap)
            rendered = windll.user32.PrintWindow(hwnd, save_dc.GetSafeHdc(), 2)
            if rendered:
                bits = bitmap.GetBitmapBits(True)
                image = Image.frombuffer("RGB", (width, height), bits, "raw", "BGRX", 0, 1).copy()
        except Exception:
            image = None
        finally:
            try:
                if bitmap is not None:
                    win32gui.DeleteObject(bitmap.GetHandle())
                if save_dc is not None:
                    save_dc.DeleteDC()
                if window_dc is not None:
                    window_dc.DeleteDC()
                if hwnd_dc is not None:
                    win32gui.ReleaseDC(hwnd, hwnd_dc)
            except Exception:
                pass

        if image is None or image.resize((1, 1)).getextrema() == ((0, 0), (0, 0), (0, 0)):
            if win32gui.IsIconic(hwnd):
                raise RuntimeError("微信窗口已最小化，请先恢复微信主界面。")
            image = ImageGrab.grab(bbox=(left, top, right, bottom), all_screens=True)
        return image

    @staticmethod
    def _unread_row_centers(image) -> list[int]:
        width, height = image.size
        x_start = max(42, int(width * 0.035))
        x_end = min(max(320, int(width * 0.42)), 520, width)
        rgb = image.convert("RGB")
        active_rows: list[int] = []
        for y in range(4, height - 4, 2):
            red_pixels = 0
            for x in range(x_start, x_end, 3):
                red, green, blue = rgb.getpixel((x, y))
                if red >= 185 and green <= 105 and blue <= 105 and red >= green + 70:
                    red_pixels += 1
                    if red_pixels >= 2:
                        active_rows.append(y)
                        break
        groups: list[list[int]] = []
        for y in active_rows:
            if not groups or y - groups[-1][-1] > 10:
                groups.append([y])
            else:
                groups[-1].append(y)
        centers = [int(sum(group) / len(group)) for group in groups if len(group) >= 2]
        return [center for center in centers if 45 <= center <= height - 35]

    @staticmethod
    def _ocr_text(image) -> str:
        try:
            import image_ocr
        except ImportError as exc:
            raise RuntimeError("未找到 Windows OCR 模块。") from exc
        if not image_ocr.ocr_available():
            raise RuntimeError(image_ocr.unavailable_reason())
        words = image_ocr.recognize_words(image)
        return image_ocr.join_words(words).strip()

    @classmethod
    def _read_uia(cls, hwnd: int) -> list[tuple[str, str]]:
        try:
            from pywinauto import Desktop  # type: ignore
        except ImportError:
            return []
        result: list[tuple[str, str]] = []
        try:
            root = Desktop(backend="uia").window(handle=hwnd)
            for control in root.descendants(control_type="ListItem"):
                text = re.sub(r"\s+", " ", str(control.window_text() or "")).strip()
                if not text or not re.search(r"未读|新消息|\d+\s*条消息", text):
                    continue
                sender, _, body = text.partition(" ")
                result.append((sender[:160] or "微信会话", body.strip() or text))
        except Exception:
            return []
        return result

    @classmethod
    def _read_ocr(cls, image) -> list[tuple[str, str]]:
        width, height = image.size
        list_width = min(max(330, int(width * 0.43)), 540, width)
        result: list[tuple[str, str]] = []
        for center in cls._unread_row_centers(image):
            top, bottom = max(0, center - 45), min(height, center + 48)
            text = cls._ocr_text(image.crop((40, top, list_width, bottom)))
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            if not text:
                continue
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            sender = next((line for line in lines if not re.fullmatch(r"\d+", line)), "微信会话")
            result.append((sender[:160], text[:4000]))
        return result

    def read(self) -> list[WeChatNotification]:
        if not self._running:
            return []
        hwnd = self._find_window()
        if not hwnd:
            raise RuntimeError("微信主窗口不可见，请打开微信并保持窗口未最小化。")
        self.last_window_handle = hwnd
        rows = self._read_uia(hwnd)
        if not rows:
            rows = self._read_ocr(self._capture_window(hwnd))
        self.last_capture_at = datetime.now().isoformat(timespec="seconds")
        self.last_unread_count = len(rows)
        result: list[WeChatNotification] = []
        for sender, text in rows:
            stable_text = re.sub(r"(?<!\d)\d{1,2}:\d{2}(?!\d)", "", text)
            raw_key = f"{hwnd}\n{sender}\n{stable_text}"
            fingerprint = hashlib.sha256(raw_key.encode("utf-8", errors="replace")).hexdigest()
            notification_id = int(fingerprint[:12], 16)
            app_id = "wechat-window-group" if any(term in sender for term in ("群", "班")) else "wechat-window"
            result.append(WeChatNotification(
                notification_id, app_id, sender, text, self.last_capture_at, fingerprint,
            ))
        return result


class AiraUsageReminder:
    """Track active computer time without collecting application or input content."""

    def __init__(
        self,
        app,
        data_dir: Path,
        *,
        clock: Callable[[], float] = time.monotonic,
        wall_clock: Callable[[], datetime] = datetime.now,
        idle_provider: Callable[[], float] = windows_idle_seconds,
    ) -> None:
        self.app = app
        self.data_dir = Path(data_dir)
        self.stats_file = self.data_dir / "usage.json"
        self._clock = clock
        self._wall_clock = wall_clock
        self._idle_provider = idle_provider
        self._stop_event = threading.Event()
        self._thread: threading.Thread | None = None
        self._lock = threading.Lock()
        self.running = False
        self.day = self._wall_clock().date().isoformat()
        self.today_seconds = 0.0
        self.continuous_seconds = 0.0
        self.last_idle_seconds = 0.0
        self.reminded_this_session = False
        self._last_tick = self._clock()
        self._last_save_tick = self._last_tick
        self._load_stats()

    def _report_unexpected(self, exc: BaseException, action: str) -> None:
        reporter = getattr(self.app, "_log_unexpected", None)
        if not callable(reporter):
            return
        try:
            reporter(
                exc,
                module="aira.usage_reminder",
                action=action,
                target_path=self.stats_file,
                expected=(OSError, ValueError, TypeError, json.JSONDecodeError),
            )
        except Exception:
            pass

    def _load_stats(self) -> None:
        try:
            raw = json.loads(self.stats_file.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError, json.JSONDecodeError):
            return
        except Exception as exc:
            self._report_unexpected(exc, "load_stats")
            return
        if not isinstance(raw, dict) or str(raw.get("day") or "") != self.day:
            return
        try:
            today_seconds = max(0.0, float(raw.get("today_seconds") or 0.0))
            continuous_seconds = max(0.0, float(raw.get("continuous_seconds") or 0.0))
            updated_at = datetime.fromisoformat(str(raw.get("updated_at") or ""))
            gap = max(0.0, (self._wall_clock() - updated_at).total_seconds())
        except (TypeError, ValueError, OverflowError):
            return
        self.today_seconds = today_seconds
        if gap < USAGE_IDLE_RESET_SECONDS:
            self.continuous_seconds = continuous_seconds
            self.reminded_this_session = bool(raw.get("reminded_this_session", False))

    def _save_stats(self) -> None:
        with self._lock:
            payload = {
                "schema_version": 1,
                "day": self.day,
                "today_seconds": int(self.today_seconds),
                "continuous_seconds": int(self.continuous_seconds),
                "reminded_this_session": bool(self.reminded_this_session),
                "updated_at": self._wall_clock().isoformat(timespec="seconds"),
            }
        try:
            self.data_dir.mkdir(parents=True, exist_ok=True)
            temporary = self.stats_file.with_suffix(".tmp")
            temporary.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8",
            )
            os.replace(temporary, self.stats_file)
        except OSError:
            return
        except Exception as exc:
            self._report_unexpected(exc, "save_stats")

    def _post(self, callback: Callable[[], None]) -> None:
        try:
            self.app.root.after(0, callback)
        except Exception:
            pass

    def _publish_status(self) -> None:
        def apply() -> None:
            window = getattr(self.app, "aira_window", None)
            if window is not None and not getattr(window, "closed", True):
                window.refresh_usage_status()

        self._post(apply)

    def _roll_day_if_needed(self) -> None:
        current_day = self._wall_clock().date().isoformat()
        if current_day == self.day:
            return
        self.day = current_day
        self.today_seconds = 0.0
        self.continuous_seconds = 0.0
        self.reminded_this_session = False

    def _send_reminder(self) -> None:
        mode = normalize_aira_notify_mode(
            self.app.settings.get("aira_usage_notify_mode", "windows")
        )
        title = "Aira 使用时长提醒"
        message = (
            f"你已连续使用电脑约 {_format_usage_duration(self.continuous_seconds)}，"
            f"今天累计 {_format_usage_duration(self.today_seconds)}。"
            "建议起身活动、远眺并休息几分钟。"
        )

        def apply() -> None:
            callback = getattr(self.app, "notify_aira_usage_reminder", None)
            if callable(callback):
                callback(title, message, mode)
                return
            status = getattr(self.app, "write_status", None)
            if callable(status):
                status(f"{title}：{message}")

        self._post(apply)

    def record_sample(self, elapsed_seconds: float, idle_seconds: float) -> dict:
        """Record one aggregate sample; exposed to keep the policy deterministic in tests."""
        elapsed = max(0.0, float(elapsed_seconds))
        idle = max(0.0, float(idle_seconds))
        should_remind = False
        with self._lock:
            self._roll_day_if_needed()
            self.last_idle_seconds = idle
            if idle >= USAGE_IDLE_RESET_SECONDS:
                self.continuous_seconds = 0.0
                self.reminded_this_session = False
            else:
                self.today_seconds += elapsed
                self.continuous_seconds += elapsed
                if (
                    self.continuous_seconds >= USAGE_REMIND_AFTER_SECONDS
                    and not self.reminded_this_session
                ):
                    self.reminded_this_session = True
                    should_remind = True
        if should_remind:
            self._send_reminder()
        self._publish_status()
        return self.snapshot()

    def _sample_once(self) -> None:
        tick = self._clock()
        elapsed = max(0.0, tick - self._last_tick)
        self._last_tick = tick
        # Do not count long suspend/resume gaps as active computer use.
        elapsed = min(elapsed, max(60.0, USAGE_POLL_SECONDS * 3.0))
        try:
            idle = self._idle_provider()
        except OSError:
            return
        except Exception as exc:
            self._report_unexpected(exc, "read_idle_time")
            return
        self.record_sample(elapsed, idle)
        if tick - self._last_save_tick >= USAGE_SAVE_INTERVAL_SECONDS:
            self._last_save_tick = tick
            self._save_stats()

    def _run(self) -> None:
        while not self._stop_event.wait(USAGE_POLL_SECONDS):
            if not self.running:
                break
            try:
                self._sample_once()
            except Exception as exc:
                self._report_unexpected(exc, "sample_usage")

    def start(self) -> None:
        if self.running:
            return
        self.running = True
        self._stop_event.clear()
        self._last_tick = self._clock()
        self._last_save_tick = self._last_tick
        self._thread = threading.Thread(
            target=self._run, daemon=True, name="Passer-Aira-Usage",
        )
        self._thread.start()
        self._publish_status()

    def stop(self) -> None:
        was_running = self.running
        self.running = False
        self._stop_event.set()
        thread = self._thread
        self._thread = None
        if thread is not None and thread is not threading.current_thread():
            thread.join(timeout=1.0)
        if was_running:
            self._save_stats()
        self._publish_status()

    close = stop

    def snapshot(self) -> dict:
        with self._lock:
            return {
                "running": bool(self.running),
                "day": self.day,
                "today_seconds": int(self.today_seconds),
                "continuous_seconds": int(self.continuous_seconds),
                "idle_seconds": int(self.last_idle_seconds),
                "privacy": "仅统计键鼠活跃时长，不记录应用、窗口标题或输入内容。",
            }

    def status_text(self) -> str:
        state = self.snapshot()
        prefix = "提醒已开启" if state["running"] else "提醒未开启"
        return (
            f"{prefix} · 今日 {_format_usage_duration(state['today_seconds'])} · "
            f"连续 {_format_usage_duration(state['continuous_seconds'])}"
        )


class AiraService:
    def __init__(self, app) -> None:
        self.app = app
        self.reader = WeChatWindowNotificationReader(app)
        self.running = False
        self.status = "Aira 监听未开启。"
        self.last_error = ""
        self._stop_event = threading.Event()
        self._wake_event = threading.Event()
        self._thread: threading.Thread | None = None
        self._seen: set[str] = set()
        self._history_lock = threading.Lock()
        self.history = self._load_history()
        self.usage_reminder = AiraUsageReminder(app, self.data_dir)
        if bool(self.app.settings.get("aira_usage_reminder_enabled", False)):
            self.usage_reminder.start()

    @property
    def data_dir(self) -> Path:
        return Path(getattr(self.app, "data_dir", Path.home() / "PasserData")) / "Aira"

    @property
    def history_file(self) -> Path:
        return self.data_dir / "history.json"

    def _load_history(self) -> list[dict]:
        try:
            raw = json.loads(self.history_file.read_text(encoding="utf-8"))
        except Exception:
            raw = []
        return [dict(item) for item in raw if isinstance(item, dict)][-HISTORY_LIMIT:]

    def _save_history(self) -> None:
        self.data_dir.mkdir(parents=True, exist_ok=True)
        temporary = self.history_file.with_suffix(".tmp")
        with self._history_lock:
            payload = list(self.history[-HISTORY_LIMIT:])
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temporary, self.history_file)

    def _post(self, callback: Callable[[], None]) -> None:
        try:
            self.app.root.after(0, callback)
        except Exception:
            pass

    def _publish_status(self, text: str, error: str = "") -> None:
        self.status = str(text)
        self.last_error = str(error)

        def apply() -> None:
            window = getattr(self.app, "aira_window", None)
            if window is not None and not getattr(window, "closed", True):
                window.refresh_status()

        self._post(apply)

    def _disable_saved_monitor(self) -> None:
        def apply() -> None:
            self.app.settings["aira_monitor_enabled"] = False
            try:
                self.app.save()
            except Exception:
                pass

        self._post(apply)

    def _interval(self) -> float:
        try:
            return float(max(2, min(30, int(self.app.settings.get("aira_poll_interval") or 3))))
        except (TypeError, ValueError):
            return 3.0

    def _contact_allowed(self, sender: str) -> bool:
        raw = str(self.app.settings.get("aira_contacts") or "").strip()
        if not raw:
            return True
        names = [part.strip().casefold() for part in re.split(r"[,，、;；\n]+", raw) if part.strip()]
        folded = str(sender or "").casefold()
        return any(name in folded for name in names)

    def start(self) -> None:
        if self.running:
            return
        self.running = True
        self._stop_event.clear()
        self._wake_event.clear()
        self._publish_status("Aira 正在连接桌面微信窗口……")
        self._thread = threading.Thread(target=self._run, daemon=True, name="Passer-Aira")
        self._thread.start()

    def stop(self) -> None:
        self.running = False
        self._stop_event.set()
        self._wake_event.set()
        self.reader.stop_channel()
        self._publish_status("Aira 监听已暂停。")

    @property
    def usage_running(self) -> bool:
        return bool(self.usage_reminder.running)

    def start_usage_reminder(self) -> None:
        self.usage_reminder.start()

    def stop_usage_reminder(self) -> None:
        self.usage_reminder.stop()

    def usage_snapshot(self) -> dict:
        return self.usage_reminder.snapshot()

    def usage_status_text(self) -> str:
        return self.usage_reminder.status_text()

    def close(self) -> None:
        self.running = False
        self._stop_event.set()
        self._wake_event.set()
        self.reader.stop_channel()
        self.usage_reminder.close()

    def check_now(self) -> None:
        if self.running:
            self._wake_event.set()
            return
        threading.Thread(target=self._check_once, daemon=True, name="Passer-Aira-Check").start()

    def _check_once(self) -> None:
        access = self.reader.ensure_access()
        if access != "allowed":
            detail = {
                "wechat_not_running": "请先启动并登录桌面微信。",
                "window_hidden": "请打开微信主界面并保持窗口未最小化。",
            }.get(access, access)
            self._publish_status("Aira 无法读取微信窗口。", detail)
            return
        try:
            self.reader.start_channel()
            notifications = self.reader.read()
            fresh = [
                item for item in notifications
                if item.fingerprint not in self._seen and self._contact_allowed(item.sender)
            ]
            self._seen.update(item.fingerprint for item in notifications)
            if fresh:
                self._summarize(fresh)
            else:
                self._publish_status("当前没有检测到新的微信未读会话。")
        except Exception as exc:
            self._publish_status("Aira 检查消息失败。", str(exc))
        finally:
            self.reader.stop_channel()

    def clear_history(self) -> None:
        with self._history_lock:
            self.history.clear()
        self._save_history()
        window = getattr(self.app, "aira_window", None)
        if window is not None and not getattr(window, "closed", True):
            window.refresh_history()

    def _run(self) -> None:
        first_poll = True
        while self.running and not self._stop_event.is_set():
            if not first_poll:
                self._wake_event.wait(self._interval())
                self._wake_event.clear()
                if not self.running or self._stop_event.is_set():
                    break
            first_poll = False

            access = self.reader.ensure_access()
            if access != "allowed":
                self.reader.stop_channel()
                detail = {
                    "wechat_not_running": "请先启动并登录桌面微信；监听会在微信出现后自动继续。",
                    "window_hidden": "请打开或恢复微信主界面；监听会自动继续。",
                }.get(access, access)
                self._publish_status("Aira 正在等待微信窗口。", detail)
                continue
            try:
                if not self.reader._running:
                    self.reader.start_channel()
                notifications = self.reader.read()
                fresh = [
                    item for item in notifications
                    if item.fingerprint not in self._seen and self._contact_allowed(item.sender)
                ]
                self._seen.update(item.fingerprint for item in notifications)
                if len(self._seen) > 2000:
                    self._seen = {item.fingerprint for item in notifications}
                if fresh:
                    self._summarize(fresh)
                else:
                    self._publish_status("Aira 正在读取微信窗口中的新未读会话。")
            except Exception as exc:
                if not self.running or self._stop_event.is_set():
                    break
                self.reader.stop_channel()
                self._publish_status("Aira 检查消息失败，稍后重试。", str(exc))
        self.reader.stop_channel()

    def _summarize(self, notifications: list[WeChatNotification]) -> None:
        self._publish_status(f"Aira 正在总结 {len(notifications)} 条微信消息……")
        try:
            analysis = self.app.summarize_aira_notifications(notifications)
            if isinstance(analysis, dict):
                summary = str(analysis.get("summary") or "").strip()
                excel_preview = analysis.get("excel_preview")
            else:
                summary = str(analysis or "").strip()
                excel_preview = None
            error = ""
        except Exception as exc:
            summary = "总结失败：" + str(exc)
            excel_preview = None
            error = str(exc)
        senders = list(dict.fromkeys(item.sender for item in notifications if item.sender))
        record = {
            "id": hashlib.sha256("|".join(item.fingerprint for item in notifications).encode("ascii")).hexdigest()[:20],
            "time": datetime.now().isoformat(timespec="seconds"),
            "senders": senders,
            "count": len(notifications),
            "summary": str(summary).strip()[:6000],
            "error": error,
            "notify_mode": normalize_aira_notify_mode(
                self.app.settings.get("aira_notify_mode", "windows")
            ),
        }
        if isinstance(excel_preview, dict):
            record["excel_preview"] = excel_preview
        if bool(self.app.settings.get("aira_save_raw", False)):
            record["messages"] = [
                {"sender": item.sender, "text": item.text, "time": item.created_at}
                for item in notifications
            ]
        with self._history_lock:
            self.history.append(record)
            self.history = self.history[-HISTORY_LIMIT:]
        try:
            self._save_history()
        except OSError:
            pass

        def apply() -> None:
            window = getattr(self.app, "aira_window", None)
            if window is not None and not getattr(window, "closed", True):
                window.refresh_history(select_id=record["id"])
            self.app.write_status(f"Aira 已总结 {len(notifications)} 条微信新消息。")
            try:
                self.app.notify_aira_summary(record)
            except Exception:
                pass

        self._post(apply)
        self._publish_status("Aira 正在读取微信窗口中的新未读会话。")


class PasserDropdown(tk.Frame):
    """Passer-styled readonly selector with a themed popup menu."""

    def __init__(self, parent, variable: tk.StringVar, values, theme: ClickerTheme,
                 font, command: Callable[[], None] | None = None,
                 width: int = 166, height: int = 38):
        super().__init__(
            parent, bg="#ffffff", highlightthickness=1,
            highlightbackground=theme.border, highlightcolor=theme.accent,
            width=width, height=height, cursor="hand2", takefocus=True,
        )
        self.variable = variable
        self.values = tuple(str(value) for value in values)
        self.theme = theme
        self.command = command
        self._popup: tk.Toplevel | None = None
        self._popup_rows: list[tuple[str, tk.Label]] = []
        self.pack_propagate(False)
        self.value_label = tk.Label(
            self, textvariable=variable, bg="#ffffff", fg="#1f2937",
            anchor=tk.W, font=font, cursor="hand2",
        )
        self.value_label.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 3))
        self.arrow_label = tk.Label(
            self, text="▾", bg="#ffffff", fg=theme.muted_fg,
            font=font, cursor="hand2",
        )
        self.arrow_label.pack(side=tk.RIGHT, padx=(3, 9))
        for child in (self, self.value_label, self.arrow_label):
            child.bind("<Button-1>", self.open_menu)
            child.bind("<Enter>", self._hover_on)
            child.bind("<Leave>", self._hover_off)
        self.bind("<Return>", self.open_menu)
        self.bind("<space>", self.open_menu)
        self.bind("<Down>", self.open_menu)
        self.bind("<FocusIn>", self._focus_on)
        self.bind("<FocusOut>", self._focus_off)

    def _set_surface(self, color: str) -> None:
        self.configure(bg=color)
        self.value_label.configure(bg=color)
        self.arrow_label.configure(bg=color)

    def _hover_on(self, _event=None) -> None:
        self.configure(highlightbackground=self.theme.accent)
        self._set_surface(self.theme.accent_soft)

    def _hover_off(self, _event=None) -> None:
        if self._popup is not None:
            return
        if self.focus_get() is not self:
            self.configure(highlightbackground=self.theme.border)
        self._set_surface("#ffffff")

    def _focus_on(self, _event=None) -> None:
        self.configure(highlightbackground=self.theme.accent)

    def _focus_off(self, _event=None) -> None:
        if self._popup is not None:
            return
        self.configure(highlightbackground=self.theme.border)

    def _choose(self, value: str) -> None:
        self.close_menu()
        self.variable.set(value)
        self.event_generate("<<ComboboxSelected>>")
        if self.command is not None:
            self.command()

    def close_menu(self, _event=None):
        popup = self._popup
        self._popup = None
        self._popup_rows = []
        if popup is not None:
            try:
                popup.destroy()
            except tk.TclError:
                pass
        self.configure(highlightbackground=self.theme.border)
        self._set_surface("#ffffff")
        return "break"

    def _close_menu_if_focus_left(self) -> None:
        popup = self._popup
        if popup is None:
            return
        try:
            focus = popup.focus_get()
            if focus is not None and str(focus).startswith(str(popup)):
                return
        except tk.TclError:
            return
        self.close_menu()

    def _native_screen_rect(self) -> tuple[int, int, int, int]:
        """Return physical screen coordinates, including negative monitor origins."""
        self.update_idletasks()
        if os.name == "nt":
            class Rect(ctypes.Structure):
                _fields_ = (
                    ("left", ctypes.c_long), ("top", ctypes.c_long),
                    ("right", ctypes.c_long), ("bottom", ctypes.c_long),
                )

            rect = Rect()
            try:
                if ctypes.windll.user32.GetWindowRect(
                    ctypes.c_void_p(int(self.winfo_id())), ctypes.byref(rect)
                ):
                    return int(rect.left), int(rect.top), int(rect.right), int(rect.bottom)
            except (AttributeError, OSError, TypeError, ValueError, tk.TclError):
                pass
        left = int(self.winfo_rootx())
        top = int(self.winfo_rooty())
        return left, top, left + int(self.winfo_width()), top + int(self.winfo_height())

    @staticmethod
    def _virtual_screen_rect(popup: tk.Toplevel) -> tuple[int, int, int, int]:
        if os.name == "nt":
            try:
                user32 = ctypes.windll.user32
                left = int(user32.GetSystemMetrics(76))
                top = int(user32.GetSystemMetrics(77))
                width = max(1, int(user32.GetSystemMetrics(78)))
                height = max(1, int(user32.GetSystemMetrics(79)))
                return left, top, left + width, top + height
            except (AttributeError, OSError, TypeError, ValueError):
                pass
        return 0, 0, int(popup.winfo_screenwidth()), int(popup.winfo_screenheight())

    def open_menu(self, _event=None):
        if not self.values:
            return "break"
        if self._popup is not None:
            return self.close_menu()
        self.focus_set()
        self.configure(highlightbackground=self.theme.accent)
        self._set_surface(self.theme.accent_soft)

        popup = tk.Toplevel(self)
        self._popup = popup
        popup.withdraw()
        popup.overrideredirect(True)
        popup.configure(bg=self.theme.border)
        owner = self.winfo_toplevel()
        try:
            popup.transient(owner)
            popup.attributes("-topmost", True)
        except tk.TclError:
            pass

        panel = tk.Frame(
            popup, bg="#ffffff", highlightthickness=1,
            highlightbackground=self.theme.border,
        )
        panel.pack(fill=tk.BOTH, expand=True)
        current = self.variable.get()
        for value in self.values:
            selected = value == current
            row = tk.Label(
                panel,
                text=(f"✓  {value}" if selected else f"    {value}"),
                bg=(self.theme.accent_soft if selected else "#ffffff"),
                fg=(self.theme.accent if selected else "#1f2937"),
                anchor=tk.W,
                padx=10,
                pady=7,
                font=self.value_label.cget("font"),
                cursor="hand2",
            )
            row.pack(fill=tk.X)
            row.bind(
                "<Enter>",
                lambda _event, widget=row: widget.configure(
                    bg=self.theme.accent, fg="#ffffff",
                ),
            )
            row.bind(
                "<Leave>",
                lambda _event, widget=row, item=value: widget.configure(
                    bg=(self.theme.accent_soft if self.variable.get() == item else "#ffffff"),
                    fg=(self.theme.accent if self.variable.get() == item else "#1f2937"),
                ),
            )
            row.bind(
                "<Button-1>",
                lambda _event, item=value: self._choose(item),
            )
            self._popup_rows.append((value, row))

        popup.update_idletasks()
        control_left, control_top, control_right, control_bottom = self._native_screen_rect()
        popup_width = max(control_right - control_left, popup.winfo_reqwidth())
        popup_height = popup.winfo_reqheight()
        screen_left, screen_top, screen_right, screen_bottom = self._virtual_screen_rect(popup)
        x = max(screen_left, min(control_left, screen_right - popup_width))
        y = control_bottom
        if y + popup_height > screen_bottom:
            y = max(screen_top, control_top - popup_height)
        self.theme.place_toplevel_absolute(
            popup, popup_width, popup_height, x, y,
        )
        popup.bind("<Escape>", self.close_menu)
        popup.bind(
            "<FocusOut>",
            lambda _event: popup.after(20, self._close_menu_if_focus_left),
        )
        try:
            popup.deiconify()
            popup.lift()
            popup.focus_force()
        except tk.TclError:
            self.close_menu()
        return "break"


class AiraMemoryDirectoryWindow:
    WIDTH = 760
    HEIGHT = 500
    TILE_WIDTH = 104
    TILE_HEIGHT = 92
    TILE_GAP = 10
    BODY_PAD = 14

    def __init__(self, owner: "AiraWindow"):
        self.owner = owner
        self.app = owner.app
        self.theme = owner.theme
        self.closed = False
        self.move_start = None
        self.columns = 0
        self.tiles: list[tuple[tk.Canvas, Path]] = []
        self.selected_path: Path | None = None
        self.status_var = tk.StringVar(value="")

        self.window = tk.Toplevel(self.app.root)
        self.window.withdraw()
        self.window.overrideredirect(True)
        self.window.configure(bg=self.theme.border)
        self.window.resizable(False, False)
        self.shell = tk.Frame(
            self.window, bg=self.theme.surface_bg,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        self.shell.pack(fill=tk.BOTH, expand=True, padx=1, pady=1)
        self._build()
        self.window.bind("<Escape>", lambda _event: self.close())
        self.window.protocol("WM_DELETE_WINDOW", self.close)

        x, y = self.theme.center_over_root(self.app.root, self.WIDTH, self.HEIGHT)
        self.theme.place_toplevel_absolute(self.window, self.WIDTH, self.HEIGHT, x, y)
        self.window.attributes("-topmost", self.app.topmost_var.get())
        self.app.apply_window_transparency(self.window)
        self.refresh()
        self.window.deiconify()
        self.app.keep_window_above_main(self.window)
        self.window.focus_set()

    def _font(self, size=9, weight="normal"):
        return self.theme.app_font(size, weight)

    def _build(self) -> None:
        bar = tk.Frame(self.shell, bg=self.theme.title_bg, height=46)
        bar.pack(fill=tk.X)
        bar.pack_propagate(False)
        title = tk.Label(
            bar, text="Aira - 记忆目录", bg=self.theme.title_bg, fg="#dbe7ff",
            anchor=tk.W, font=self._font(10, "bold"),
        )
        title.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=14)
        close = tk.Button(
            bar, text="×", command=self.close, bd=0, padx=11, pady=5,
            bg=self.theme.title_button_bg, fg="#e7eefc", activebackground="#ef4444",
            activeforeground="#ffffff", cursor="hand2", font=self._font(10),
        )
        close.pack(side=tk.RIGHT, padx=(0, 8), pady=7)
        refresh = tk.Button(
            bar, text="↻", command=self.refresh, bd=0, width=3, padx=0, pady=5,
            bg=self.theme.title_button_bg, fg="#e7eefc",
            activebackground=self.theme.title_button_hover, activeforeground="#ffffff",
            cursor="hand2", font=self._font(11),
        )
        refresh.pack(side=tk.RIGHT, padx=(0, 6), pady=7)
        for widget in (bar, title):
            widget.bind("<ButtonPress-1>", self._start_move)
            widget.bind("<B1-Motion>", self._do_move)
            widget.bind("<ButtonRelease-1>", self._finish_move)

        body = tk.Frame(self.shell, bg="#ffffff")
        body.pack(fill=tk.BOTH, expand=True)
        self.canvas = tk.Canvas(body, bg="#ffffff", bd=0, highlightthickness=0)
        self.canvas.pack(fill=tk.BOTH, expand=True)
        self.content = tk.Frame(self.canvas, bg="#ffffff")
        self.content_window = self.canvas.create_window(
            (0, 0), window=self.content, anchor=tk.NW,
        )
        self.content.bind("<Configure>", self._update_scroll_region)
        self.canvas.bind("<Configure>", self._on_canvas_configure)
        self.canvas.bind("<MouseWheel>", self._on_mouse_wheel)

        status = tk.Frame(self.shell, bg=self.theme.title_bg, height=24)
        status.pack(fill=tk.X)
        status.pack_propagate(False)
        tk.Label(
            status, textvariable=self.status_var, bg=self.theme.title_bg,
            fg="#9fb3d4", anchor=tk.W, font=self._font(8),
        ).pack(fill=tk.BOTH, expand=True, padx=10)

    def _memory_files(self) -> list[Path]:
        try:
            files = [path for path in self.owner.memory_dir.glob("*.md") if path.is_file()]
        except OSError:
            return []

        def sort_key(path: Path):
            try:
                modified = path.stat().st_mtime_ns
            except OSError:
                modified = 0
            return (-modified, path.name.casefold())

        return sorted(files, key=sort_key)

    def _display_name(self, path: Path) -> str:
        name = path.name
        if len(name) <= 24:
            return name
        return f"{name[:20]}…{path.suffix}"

    def _create_tile(self, path: Path) -> tk.Canvas:
        tile = tk.Canvas(
            self.content, width=self.TILE_WIDTH, height=self.TILE_HEIGHT,
            bg="#ffffff", bd=0, highlightthickness=0, cursor="hand2",
        )
        center = self.TILE_WIDTH // 2
        left, top, right, bottom = center - 15, 8, center + 15, 46
        tile.create_polygon(
            left, top, right - 8, top, right, top + 8, right, bottom,
            left, bottom, left, top,
            fill="#ffffff", outline=self.theme.border, width=1,
        )
        tile.create_line(right - 8, top, right - 8, top + 8, right, top + 8,
                         fill=self.theme.border, width=1)
        tile.create_rectangle(
            left + 4, bottom - 15, right - 4, bottom - 4,
            fill=self.theme.accent, outline=self.theme.accent,
        )
        tile.create_text(
            center, bottom - 9, text="MD", fill="#ffffff",
            font=self._font(6, "bold"),
        )
        tile.create_text(
            center, 66, text=self._display_name(path), fill="#334155",
            width=self.TILE_WIDTH - 10, justify=tk.CENTER,
            font=self._font(8),
        )
        tile.bind("<Button-1>", lambda _event, p=path: self._select(p))
        tile.bind("<Double-Button-1>", lambda _event, p=path: self._open(p))
        tile.bind("<Return>", lambda _event, p=path: self._open(p))
        tile.bind("<MouseWheel>", self._on_mouse_wheel)
        return tile

    def _select(self, path: Path) -> None:
        self.selected_path = path
        for tile, tile_path in self.tiles:
            tile.configure(bg=self.theme.accent_soft if tile_path == path else "#ffffff")
        self.status_var.set(f"已选择：{path.name}（双击编辑）")

    def _open(self, path: Path) -> None:
        self.close()
        self.owner.open_memory_file(path)

    def refresh(self) -> None:
        try:
            self.owner.memory_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            self.status_var.set(f"无法读取记忆目录：{exc}")
            return
        for child in self.content.winfo_children():
            child.destroy()
        self.tiles.clear()
        self.selected_path = None
        files = self._memory_files()
        if not files:
            tk.Label(
                self.content, text="暂无记忆文件", bg="#ffffff", fg="#94a3b8",
                anchor=tk.W, font=self._font(9),
            ).grid(row=0, column=0, sticky="nw", padx=20, pady=20)
        else:
            for path in files:
                self.tiles.append((self._create_tile(path), path))
        self.status_var.set(f"{len(files)} 个 Markdown 记忆文件 · 双击图标编辑")
        self._layout_tiles(force=True)
        self.window.after_idle(self._update_scroll_region)

    def _layout_tiles(self, force: bool = False) -> None:
        width = max(1, self.canvas.winfo_width())
        available = max(self.TILE_WIDTH, width - self.BODY_PAD * 2)
        columns = max(1, (available + self.TILE_GAP) // (self.TILE_WIDTH + self.TILE_GAP))
        if not force and columns == self.columns:
            return
        self.columns = columns
        for index, (tile, _path) in enumerate(self.tiles):
            row, column = divmod(index, columns)
            tile.grid(
                row=row, column=column, sticky="nw",
                padx=(self.BODY_PAD if column == 0 else self.TILE_GAP, 0),
                pady=(self.BODY_PAD if row == 0 else self.TILE_GAP, 0),
            )
        self._update_scroll_region()

    def _on_canvas_configure(self, event) -> None:
        self.canvas.itemconfigure(self.content_window, width=event.width)
        self._layout_tiles()

    def _update_scroll_region(self, _event=None) -> None:
        try:
            self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        except tk.TclError:
            pass

    def _on_mouse_wheel(self, event) -> str:
        self.canvas.yview_scroll(-int(event.delta / 120), "units")
        return "break"

    def _start_move(self, event) -> None:
        self.move_start = (event.x_root, event.y_root, self.window.winfo_x(), self.window.winfo_y())

    def _do_move(self, event) -> None:
        if self.move_start:
            sx, sy, wx, wy = self.move_start
            self.window.geometry(f"+{wx + event.x_root - sx}+{wy + event.y_root - sy}")

    def _finish_move(self, _event=None) -> None:
        self.move_start = None

    def show(self) -> None:
        if self.closed:
            return
        self.refresh()
        self.window.deiconify()
        self.window.lift()
        self.app.keep_window_above_main(self.window)
        self.window.focus_set()

    def close(self) -> None:
        if self.closed:
            return
        self.closed = True
        try:
            self.window.destroy()
        except Exception:
            pass
        if getattr(self.owner, "memory_directory_window", None) is self:
            self.owner.memory_directory_window = None


class AiraMemoryEditorWindow:
    WIDTH = 760
    HEIGHT = 560

    def __init__(self, owner: "AiraWindow", path: Path):
        self.owner = owner
        self.app = owner.app
        self.theme = owner.theme
        self.path = path
        self.closed = False
        self.move_start = None
        self.original = ""
        self.status_var = tk.StringVar(value=str(path))

        self.window = tk.Toplevel(self.app.root)
        self.window.withdraw()
        self.window.overrideredirect(True)
        self.window.configure(bg=self.theme.border)
        self.window.resizable(False, False)
        self.shell = tk.Frame(
            self.window, bg=self.theme.surface_bg,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        self.shell.pack(fill=tk.BOTH, expand=True, padx=1, pady=1)
        self._build()
        self._load()
        self.window.bind("<Escape>", lambda _event: self.close())
        self.window.bind("<Control-s>", self.save)
        self.window.bind("<Control-S>", self.save)
        self.window.protocol("WM_DELETE_WINDOW", self.close)

        x, y = self.theme.center_over_root(self.app.root, self.WIDTH, self.HEIGHT)
        self.theme.place_toplevel_absolute(self.window, self.WIDTH, self.HEIGHT, x, y)
        self.window.attributes("-topmost", self.app.topmost_var.get())
        self.app.apply_window_transparency(self.window)
        self.window.deiconify()
        self.app.keep_window_above_main(self.window)
        try:
            self.app.focus_manager.claim(self.window, self.editor)
        except Exception:
            self.editor.focus_set()

    def _font(self, size=9, weight="normal"):
        return self.theme.app_font(size, weight)

    def _build(self) -> None:
        bar = tk.Frame(self.shell, bg=self.theme.title_bg, height=46)
        bar.pack(fill=tk.X)
        bar.pack_propagate(False)
        title = tk.Label(
            bar, text=self.path.name, bg=self.theme.title_bg, fg="#dbe7ff",
            anchor=tk.W, font=self._font(10, "bold"),
        )
        title.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=14)
        close = tk.Button(
            bar, text="×", command=self.close, bd=0, padx=11, pady=5,
            bg=self.theme.title_button_bg, fg="#e7eefc", activebackground="#ef4444",
            activeforeground="#ffffff", cursor="hand2", font=self._font(10),
        )
        close.pack(side=tk.RIGHT, padx=(0, 8), pady=7)
        for widget in (bar, title):
            widget.bind("<ButtonPress-1>", self._start_move)
            widget.bind("<B1-Motion>", self._do_move)
            widget.bind("<ButtonRelease-1>", self._finish_move)

        editor_shell = tk.Frame(
            self.shell, bg="#ffffff", highlightthickness=1,
            highlightbackground=self.theme.border,
        )
        editor_shell.pack(fill=tk.BOTH, expand=True, padx=16, pady=(16, 10))
        self.editor = tk.Text(
            editor_shell, bd=0, relief=tk.FLAT, wrap=tk.WORD, undo=True,
            bg="#ffffff", fg="#1f2937", insertbackground="#1f2937",
            selectbackground=self.theme.accent, selectforeground="#ffffff",
            highlightthickness=0, font=self._font(10), padx=12, pady=12,
        )
        self.editor.pack(fill=tk.BOTH, expand=True)
        self.editor.bind(
            "<ButtonPress-1>",
            lambda _event: self.app.focus_manager.claim(self.window, self.editor),
            add="+",
        )

        bottom = tk.Frame(self.shell, bg=self.theme.surface_bg, height=50)
        bottom.pack(fill=tk.X, padx=16, pady=(0, 12))
        bottom.pack_propagate(False)
        tk.Label(
            bottom, textvariable=self.status_var, bg=self.theme.surface_bg,
            fg="#64748b", anchor=tk.W, font=self._font(8),
        ).pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        tk.Button(
            bottom, text="保存", command=self.save, bd=0, padx=18, pady=7,
            bg=self.theme.accent, fg="#ffffff",
            activebackground=self.theme.accent_hover, activeforeground="#ffffff",
            cursor="hand2", font=self._font(9, "bold"),
        ).pack(side=tk.RIGHT, pady=8)

    def _load(self) -> None:
        try:
            self.original = self.path.read_text(encoding="utf-8-sig")
        except OSError as exc:
            self.original = ""
            self.status_var.set(f"读取失败：{exc}")
        self.editor.delete("1.0", tk.END)
        self.editor.insert("1.0", self.original)
        self.editor.edit_reset()

    def save(self, _event=None) -> str:
        if self.closed:
            return "break"
        text = self.editor.get("1.0", "end-1c")
        if text == self.original:
            self.status_var.set(f"已保存 · {self.path.name}")
            return "break"
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            temporary = self.path.with_suffix(self.path.suffix + ".tmp")
            temporary.write_text(text.rstrip() + ("\n" if text else ""), encoding="utf-8")
            os.replace(temporary, self.path)
            self.original = text
            self.status_var.set(f"已保存 · {self.path.name}")
            self.owner.refresh_memory()
        except OSError as exc:
            self.status_var.set(f"保存失败：{exc}")
            messagebox.showerror("记忆保存失败", str(exc), parent=self.window)
        return "break"

    def show(self) -> None:
        if self.closed:
            return
        self.window.deiconify()
        self.window.lift()
        self.app.keep_window_above_main(self.window)
        try:
            self.app.focus_manager.claim(self.window, self.editor)
        except Exception:
            self.editor.focus_set()

    def _start_move(self, event) -> None:
        self.move_start = (event.x_root, event.y_root, self.window.winfo_x(), self.window.winfo_y())

    def _do_move(self, event) -> None:
        if self.move_start:
            sx, sy, wx, wy = self.move_start
            self.window.geometry(f"+{wx + event.x_root - sx}+{wy + event.y_root - sy}")

    def _finish_move(self, _event=None) -> None:
        self.move_start = None

    def close(self) -> None:
        if self.closed:
            return
        self.save()
        self.closed = True
        try:
            self.window.destroy()
        except Exception:
            pass
        self.owner.memory_editor_windows.pop(self.path, None)


class AiraWindow:
    WIDTH = 1180
    HEIGHT = 726

    def __init__(self, app, theme: ClickerTheme, service: AiraService):
        self.app = app
        self.theme = theme
        self.service = service
        self.closed = False
        self.move_start = None
        self.status_var = tk.StringVar(value=service.status)
        notify_mode = normalize_aira_notify_mode(app.settings.get("aira_notify_mode", "windows"))
        self.notify_mode_var = tk.StringVar(value=AIRA_NOTIFY_MODE_LABELS[notify_mode])
        usage_notify_mode = normalize_aira_notify_mode(
            app.settings.get("aira_usage_notify_mode", "windows")
        )
        self.usage_notify_mode_var = tk.StringVar(
            value=AIRA_NOTIFY_MODE_LABELS[usage_notify_mode]
        )
        self.usage_status_var = tk.StringVar(value=service.usage_status_text())
        self.openclaw_status_var = tk.StringVar()
        self._openclaw_configuring = False
        self._openclaw_config_result: bool | None = None
        self._openclaw_config_message = ""
        self._openclaw_qr_launching = False
        self._openclaw_qr_result: bool | None = None
        self._openclaw_qr_message = ""
        self.memory_notes: list[str] = []
        self.memory_count_var = tk.StringVar(value="0 条")
        self.memory_directory_window: AiraMemoryDirectoryWindow | None = None
        self.memory_editor_windows: dict[Path, AiraMemoryEditorWindow] = {}

        self.window = tk.Toplevel(app.root)
        self.window.withdraw()
        self.window.overrideredirect(True)
        self.window.configure(bg=theme.border)
        self.window.minsize(self.WIDTH, self.HEIGHT)
        self.window.resizable(False, False)
        self.shell = tk.Frame(self.window, bg=theme.surface_bg, highlightthickness=1, highlightbackground=theme.border)
        self.shell.pack(fill=tk.BOTH, expand=True, padx=1, pady=1)
        self._build()
        self.window.bind("<Escape>", lambda _event: self.close())
        self.window.protocol("WM_DELETE_WINDOW", self.close)
        x, y = theme.center_over_root(app.root, self.WIDTH, self.HEIGHT)
        theme.place_toplevel_absolute(self.window, self.WIDTH, self.HEIGHT, x, y)
        self.window.attributes("-topmost", app.topmost_var.get())
        app.apply_window_transparency(self.window)
        self.window.deiconify()
        self.window.focus_set()
        self.refresh_memory()
        self.refresh_status()

    def _font(self, size=9, weight="normal"):
        return self.theme.app_font(size, weight)

    def _button(self, parent, text, command, primary=False, danger=False):
        bg = self.theme.danger if danger else (self.theme.accent if primary else "#eef2f7")
        fg = "#ffffff" if primary or danger else "#334155"
        return tk.Button(parent, text=text, command=command, bd=0, padx=13, pady=7,
                         bg=bg, fg=fg, activebackground=self.theme.accent_hover if primary else "#e2e8f0",
                         activeforeground="#ffffff" if primary or danger else "#111827",
                         cursor="hand2", font=self._font(9, "bold" if primary else "normal"))

    def _build(self) -> None:
        bar = tk.Frame(self.shell, bg=self.theme.title_bg, height=46)
        bar.pack(fill=tk.X)
        bar.pack_propagate(False)
        title = tk.Label(bar, text="Aira", bg=self.theme.title_bg, fg="#dbe7ff",
                         anchor=tk.W, font=self._font(11, "bold"))
        title.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=14)
        close = tk.Button(bar, text="×", command=self.close, bd=0, padx=11, pady=5,
                          bg=self.theme.title_button_bg, fg="#e7eefc", activebackground="#ef4444",
                          activeforeground="#ffffff", cursor="hand2", font=self._font(10))
        close.pack(side=tk.RIGHT, padx=(0, 8), pady=7)
        close.bind("<Enter>", lambda _event: close.configure(bg="#ef4444", fg="#ffffff"))
        close.bind(
            "<Leave>",
            lambda _event: close.configure(bg=self.theme.title_button_bg, fg="#e7eefc"),
        )
        for widget in (bar, title):
            widget.bind("<ButtonPress-1>", self._start_move)
            widget.bind("<B1-Motion>", self._do_move)
            widget.bind("<ButtonRelease-1>", self._finish_move)

        # The Aira body should read as an unpainted content surface.  Keep the
        # application background out of this area so it matches the white shell
        # instead of leaving a large tinted rectangle below the title bar.
        content = tk.Frame(self.shell, bg=self.theme.surface_bg)
        content.pack(fill=tk.BOTH, expand=True, padx=18, pady=18)

        monitor_row = tk.Frame(
            content, bg="#ffffff", height=72,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        monitor_row.pack(fill=tk.X)
        monitor_row.pack_propagate(False)
        tk.Label(
            monitor_row, text="微信监听", bg="#ffffff", fg="#0f172a",
            anchor=tk.W, font=self._font(10, "bold"),
        ).pack(side=tk.LEFT, fill=tk.Y, padx=18)
        self.monitor_button = self._button(
            monitor_row, "开启监听", self.toggle_monitor, primary=True,
        )
        self.monitor_button.pack(side=tk.RIGHT, padx=18, pady=16)
        self.notify_dropdown = PasserDropdown(
            monitor_row,
            self.notify_mode_var,
            tuple(AIRA_NOTIFY_MODE_LABELS.values()),
            self.theme,
            self._font(9),
            command=self.save_notification_mode,
        )
        self.notify_dropdown.pack(side=tk.RIGHT, padx=(0, 2), pady=17)
        self.monitor_status_label = tk.Label(
            monitor_row, textvariable=self.status_var, bg="#ffffff", fg="#64748b",
            anchor=tk.W, justify=tk.LEFT, font=self._font(8),
        )
        self.monitor_status_label.pack(
            side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(8, 10), pady=10,
        )

        memory_row = tk.Frame(
            content, bg="#ffffff", height=72,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        memory_row.pack(fill=tk.X, pady=(12, 0))
        memory_row.pack_propagate(False)
        tk.Label(
            memory_row, text="记忆", bg="#ffffff", fg="#0f172a",
            anchor=tk.W, font=self._font(10, "bold"),
        ).pack(side=tk.LEFT, fill=tk.Y, padx=(18, 10))
        tk.Label(
            memory_row, textvariable=self.memory_count_var, bg="#ffffff", fg="#64748b",
            anchor=tk.W, font=self._font(9),
        ).pack(side=tk.LEFT, fill=tk.Y)
        open_memory_dir = self._button(
            memory_row, "打开记忆目录", self.open_memory_directory, primary=True,
        )
        open_memory_dir.pack(side=tk.RIGHT, padx=18, pady=16)

        usage_row = tk.Frame(
            content, bg="#ffffff", height=72,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        usage_row.pack(fill=tk.X, pady=(12, 0))
        usage_row.pack_propagate(False)
        tk.Label(
            usage_row, text="使用时长提醒", bg="#ffffff", fg="#0f172a",
            anchor=tk.W, font=self._font(10, "bold"),
        ).pack(side=tk.LEFT, fill=tk.Y, padx=18)
        self.usage_status_label = tk.Label(
            usage_row, textvariable=self.usage_status_var, bg="#ffffff", fg="#64748b",
            anchor=tk.W, justify=tk.LEFT, font=self._font(8),
        )
        self.usage_button = self._button(
            usage_row, "开启提醒", self.toggle_usage_reminder, primary=True,
        )
        self.usage_button.pack(side=tk.RIGHT, padx=18, pady=16)
        self.usage_notify_dropdown = PasserDropdown(
            usage_row,
            self.usage_notify_mode_var,
            tuple(AIRA_NOTIFY_MODE_LABELS.values()),
            self.theme,
            self._font(9),
            command=self.save_usage_notification_mode,
        )
        self.usage_notify_dropdown.pack(side=tk.RIGHT, padx=(0, 2), pady=17)
        self.usage_status_label.pack(
            side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(8, 10), pady=10,
        )

        openclaw_row = tk.Frame(
            content, bg="#ffffff", height=72,
            highlightthickness=1, highlightbackground=self.theme.border,
        )
        openclaw_row.pack(fill=tk.X, pady=(12, 0))
        openclaw_row.pack_propagate(False)
        tk.Label(
            openclaw_row, text="OpenClaw 配置", bg="#ffffff", fg="#0f172a",
            anchor=tk.W, font=self._font(10, "bold"),
        ).pack(side=tk.LEFT, fill=tk.Y, padx=18)
        self.openclaw_status_label = tk.Label(
            openclaw_row, textvariable=self.openclaw_status_var,
            bg="#ffffff", fg="#64748b", anchor=tk.W, justify=tk.LEFT,
            font=self._font(8),
        )
        self.openclaw_button = self._button(
            openclaw_row, "自动配置", self.auto_configure_openclaw, primary=True,
        )
        self.openclaw_qr_button = self._button(
            openclaw_row, "生成微信二维码", self.generate_openclaw_wechat_qr,
            primary=True,
        )
        # Pack the QR button first so it stays at the far right; auto-configure
        # then naturally sits immediately to its left.
        self.openclaw_qr_button.pack(side=tk.RIGHT, padx=(8, 18), pady=16)
        self.openclaw_button.pack(side=tk.RIGHT, padx=(8, 0), pady=16)
        self.openclaw_status_label.pack(
            side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(8, 10), pady=10,
        )
        self.refresh_openclaw_status()

    def save_settings(self) -> None:
        try:
            self.app.save()
        except Exception:
            pass

    def save_notification_mode(self) -> None:
        mode = normalize_aira_notify_mode(self.notify_mode_var.get())
        self.notify_mode_var.set(AIRA_NOTIFY_MODE_LABELS[mode])
        self.app.settings["aira_notify_mode"] = mode
        self.app.settings["aira_notify"] = True
        self.save_settings()
        try:
            self.app.write_status(f"Aira 微信监听已改为{AIRA_NOTIFY_MODE_LABELS[mode]}。")
        except Exception:
            pass

    def save_usage_notification_mode(self) -> None:
        mode = normalize_aira_notify_mode(self.usage_notify_mode_var.get())
        self.usage_notify_mode_var.set(AIRA_NOTIFY_MODE_LABELS[mode])
        self.app.settings["aira_usage_notify_mode"] = mode
        self.save_settings()
        try:
            self.app.write_status(
                f"Aira 使用时长提醒已改为 {AIRA_NOTIFY_MODE_LABELS[mode]}。"
            )
        except Exception:
            pass

    def toggle_usage_reminder(self) -> None:
        enabled = not self.service.usage_running
        self.app.settings["aira_usage_reminder_enabled"] = enabled
        if enabled:
            self.service.start_usage_reminder()
        else:
            self.service.stop_usage_reminder()
        self.save_settings()
        self.refresh_usage_status()

    def refresh_usage_status(self) -> None:
        self.usage_status_var.set(self.service.usage_status_text())
        if self.service.usage_running:
            self.usage_status_label.configure(fg=self.theme.accent)
            self.usage_button.configure(text="关闭提醒", bg="#eef2f7", fg="#334155")
        else:
            self.usage_status_label.configure(fg="#64748b")
            self.usage_button.configure(text="开启提醒", bg=self.theme.accent, fg="#ffffff")

    @staticmethod
    def _openclaw_message_summary(message: str) -> str:
        first = next((line.strip() for line in str(message or "").splitlines() if line.strip()), "")
        return first[:110]

    def refresh_openclaw_status(self) -> None:
        qr_button = getattr(self, "openclaw_qr_button", None)
        if getattr(self, "_openclaw_configuring", False):
            self.openclaw_status_var.set("正在生成本机令牌并同步 OpenClaw MCP 配置……")
            self.openclaw_status_label.configure(fg=self.theme.accent)
            self.openclaw_button.configure(
                text="配置中…", state=tk.DISABLED, bg="#eef2f7", fg="#64748b",
            )
            if qr_button is not None:
                qr_button.configure(
                    text="生成微信二维码", state=tk.DISABLED,
                    bg="#eef2f7", fg="#64748b",
                )
            return
        if getattr(self, "_openclaw_qr_launching", False):
            self.openclaw_status_var.set("正在打开 OpenClaw 微信扫码绑定窗口……")
            self.openclaw_status_label.configure(fg=self.theme.accent)
            self.openclaw_button.configure(
                text="自动配置", state=tk.DISABLED, bg="#eef2f7", fg="#64748b",
            )
            if qr_button is not None:
                qr_button.configure(
                    text="生成中…", state=tk.DISABLED,
                    bg="#eef2f7", fg="#64748b",
                )
            return
        enabled = bool(getattr(self.app, "openclaw_enabled", False))
        qr_result = getattr(self, "_openclaw_qr_result", None)
        if qr_result is False:
            detail = self._openclaw_message_summary(self._openclaw_qr_message)
            self.openclaw_status_var.set(f"微信二维码生成失败：{detail or '请重试'}")
            color = "#ef4444"
            button_text = "自动配置" if not enabled else "重新配置"
        elif qr_result is True:
            self.openclaw_status_var.set(
                "微信二维码窗口已打开 · 请使用手机微信扫码并确认绑定"
            )
            color = self.theme.accent
            button_text = "自动配置" if not enabled else "重新配置"
        elif not enabled:
            self.openclaw_status_var.set("未启用 · 自动注册 Passer MCP；微信仍需扫码与发送者配对")
            color = "#64748b"
            button_text = "自动配置"
        elif self._openclaw_config_result is False:
            detail = self._openclaw_message_summary(self._openclaw_config_message)
            self.openclaw_status_var.set(
                f"本地桥已启用 · 自动配置未完成：{detail or '请重试'}"
            )
            color = "#ef4444"
            button_text = "重试配置"
        elif self._openclaw_config_result is True:
            self.openclaw_status_var.set(
                "配置完成 · Passer MCP 已授权；微信通道仍需扫码与发送者配对"
            )
            color = self.theme.accent
            button_text = "重新配置"
        else:
            self.openclaw_status_var.set(
                "本地桥已启用 · 可自动检查并重新注册 OpenClaw MCP"
            )
            color = self.theme.accent
            button_text = "重新配置"
        self.openclaw_status_label.configure(fg=color)
        self.openclaw_button.configure(
            text=button_text, state=tk.NORMAL, bg=self.theme.accent, fg="#ffffff",
        )
        if qr_button is not None:
            qr_button.configure(
                text="生成微信二维码", state=tk.NORMAL,
                bg=self.theme.accent, fg="#ffffff",
            )

    def auto_configure_openclaw(self) -> None:
        if self._openclaw_configuring:
            return
        self._openclaw_configuring = True
        self._openclaw_config_result = None
        self._openclaw_config_message = ""
        self._openclaw_qr_result = None
        self._openclaw_qr_message = ""
        self.refresh_openclaw_status()
        try:
            configure = getattr(self.app, "apply_openclaw_setting", None)
            if not callable(configure):
                raise RuntimeError("当前 Passer 版本不支持 OpenClaw 自动配置。")
            self.app.openclaw_enabled = True
            self.app.settings["openclaw_enabled"] = True
            self.app.save()
            configure(True, notify=False, on_complete=self._finish_openclaw_configuration)
        except (OSError, RuntimeError) as exc:
            self._finish_openclaw_configuration(False, str(exc))
        except Exception as exc:  # noqa: BLE001 - isolate the external configuration boundary
            logger = getattr(self.app, "_log_unexpected", None)
            if callable(logger):
                logger(
                    exc,
                    module="aira.openclaw",
                    action="auto_configure",
                    target_path=Path(getattr(self.app, "data_dir", Path.home())) / "OpenClaw",
                )
            self._finish_openclaw_configuration(False, str(exc))

    def _finish_openclaw_configuration(self, ok: bool, message: str) -> None:
        if self.closed:
            return
        self._openclaw_configuring = False
        self._openclaw_config_result = bool(ok)
        self._openclaw_config_message = str(message or "")
        self.refresh_openclaw_status()
        try:
            if ok:
                self.app.write_status("Aira 已自动完成 OpenClaw Passer MCP 配置。")
            else:
                self.app.write_status(
                    "Aira OpenClaw 自动配置未完成："
                    + (self._openclaw_message_summary(message) or "未知原因")
                )
        except Exception:
            pass
        if not ok:
            messagebox.showinfo(
                "OpenClaw 自动配置未完成",
                str(message or "请确认已安装 OpenClaw 后重试。"),
                parent=self.window,
            )

    def generate_openclaw_wechat_qr(self) -> None:
        if self._openclaw_configuring or self._openclaw_qr_launching:
            return
        self._openclaw_qr_launching = True
        self._openclaw_qr_result = None
        self._openclaw_qr_message = ""
        self.refresh_openclaw_status()

        def worker() -> None:
            try:
                from ai_cli_bridge import run_wechat_action

                message = str(run_wechat_action({}, "wechat_cli_login") or "")
                ok = message.startswith("已打开微信 CLI 扫码登录窗口")
            except (OSError, RuntimeError) as exc:
                ok, message = False, str(exc)
            except Exception as exc:  # noqa: BLE001 - external CLI boundary
                logger = getattr(self.app, "_log_unexpected", None)
                if callable(logger):
                    logger(
                        exc,
                        module="aira.openclaw",
                        action="generate_wechat_qr",
                        target_path=Path(getattr(self.app, "data_dir", Path.home())) / "OpenClaw",
                    )
                ok, message = False, str(exc)

            def complete() -> None:
                self._finish_openclaw_wechat_qr(ok, message)

            try:
                self.app.root.after(0, complete)
            except tk.TclError:
                pass

        threading.Thread(
            target=worker, daemon=True, name="Passer-Aira-OpenClawQR",
        ).start()

    def _finish_openclaw_wechat_qr(self, ok: bool, message: str) -> None:
        if self.closed:
            return
        self._openclaw_qr_launching = False
        self._openclaw_qr_result = bool(ok)
        self._openclaw_qr_message = str(message or "")
        self.refresh_openclaw_status()
        try:
            self.app.write_status(
                "已打开 OpenClaw 微信二维码绑定窗口。" if ok
                else "OpenClaw 微信二维码生成失败："
                + (self._openclaw_message_summary(message) or "未知原因")
            )
        except Exception:
            pass
        if not ok:
            messagebox.showinfo(
                "微信二维码生成失败",
                str(message or "请先安装并启用 OpenClaw 微信通道。"),
                parent=self.window,
            )

    def toggle_monitor(self) -> None:
        self.save_settings()
        enabled = not self.service.running
        self.app.settings["aira_monitor_enabled"] = enabled
        try:
            self.app.save()
        except Exception:
            pass
        if enabled:
            self.service.start()
        else:
            self.service.stop()
        self.refresh_status()

    def refresh_status(self) -> None:
        self.status_var.set(self.service.status + (f" {self.service.last_error}" if self.service.last_error else ""))
        if self.service.running:
            waiting = "等待微信窗口" in self.service.status
            self.monitor_status_label.configure(fg=self.theme.accent)
            self.monitor_button.configure(
                text="等待微信" if waiting else "暂停监听",
                bg="#eef2f7", fg="#334155",
            )
        else:
            self.monitor_status_label.configure(fg="#64748b")
            self.monitor_button.configure(text="开启监听", bg=self.theme.accent, fg="#ffffff")
        self.refresh_usage_status()
        self.refresh_openclaw_status()

    @property
    def memory_dir(self) -> Path:
        try:
            import ai_chat

            return Path(ai_chat.AI_MEMORY_DIR)
        except Exception:
            return self.service.data_dir / "Memory"

    def _read_memory_documents(self) -> list[tuple[Path, str]]:
        chat = getattr(self.app, "ai_chat", None)
        try:
            import ai_chat

            ai_chat.load_ai_state()  # Also migrates legacy JSON notes into Markdown files.
            documents = ai_chat.load_memory_documents()
            notes = [text for _path, text in documents]
            if chat is not None:
                chat.memory_notes[:] = notes
            return documents
        except Exception:
            return []

    def _sync_live_memory(self) -> None:
        documents = self._read_memory_documents()
        self.memory_notes = [text for _path, text in documents]
        self.memory_count_var.set(f"{len(documents)} 条")

    def save_memory(self, _event=None, *, rebuild: bool = False) -> str:
        for editor_window in list(self.memory_editor_windows.values()):
            if not editor_window.closed:
                editor_window.save()
        if rebuild:
            self.refresh_memory()
        return "break"

    def refresh_memory(self) -> None:
        documents = self._read_memory_documents()
        self.memory_count_var.set(f"{len(documents)} 条")
        self.memory_notes = [text for _path, text in documents]
        directory_window = self.memory_directory_window
        if directory_window is not None and not directory_window.closed:
            directory_window.refresh()

    def open_memory_directory(self) -> None:
        try:
            self.memory_dir.mkdir(parents=True, exist_ok=True)
            if self.memory_directory_window is not None and not self.memory_directory_window.closed:
                self.memory_directory_window.show()
            else:
                self.memory_directory_window = AiraMemoryDirectoryWindow(self)
        except Exception as exc:
            messagebox.showerror("无法打开记忆目录", str(exc), parent=self.window)

    def open_memory_file(self, path: Path) -> None:
        path = Path(path)
        if not path.exists() or not path.is_file():
            messagebox.showerror("无法编辑记忆", f"未找到记忆文件：\n{path}", parent=self.window)
            return
        editor_window = self.memory_editor_windows.get(path)
        if editor_window is not None and not editor_window.closed:
            editor_window.show()
            return
        self.memory_editor_windows[path] = AiraMemoryEditorWindow(self, path)

    def delete_memory(self, path: Path) -> None:
        try:
            editor_window = self.memory_editor_windows.get(path)
            if editor_window is not None and not editor_window.closed:
                editor_window.close()
            path.unlink(missing_ok=True)
            self.refresh_memory()
        except Exception as exc:
            messagebox.showerror("无法删除记忆", str(exc), parent=self.window)

    def clear_history(self) -> None:
        if messagebox.askyesno("清空 Aira 总结", "确定清空全部 Aira 总结记录吗？", parent=self.window):
            self.service.clear_history()

    def refresh_history(self, select_id: str = "") -> None:
        return

    def show(self) -> None:
        try:
            self.refresh_memory()
            self.refresh_status()
            self.window.deiconify(); self.window.lift(); self.window.focus_set()
        except Exception:
            pass

    def _start_move(self, event) -> None:
        self.move_start = (event.x_root, event.y_root, self.window.winfo_x(), self.window.winfo_y())

    def _do_move(self, event) -> None:
        if self.move_start:
            sx, sy, wx, wy = self.move_start
            self.window.geometry(f"+{wx + event.x_root - sx}+{wy + event.y_root - sy}")

    def _finish_move(self, _event=None) -> None:
        self.move_start = None

    def close(self) -> None:
        self.save_memory()
        self.save_settings()
        self.closed = True
        for editor_window in list(self.memory_editor_windows.values()):
            editor_window.close()
        if self.memory_directory_window is not None:
            self.memory_directory_window.close()
        try:
            self.window.destroy()
        except Exception:
            pass
        if getattr(self.app, "aira_window", None) is self:
            self.app.aira_window = None
